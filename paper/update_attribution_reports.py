"""Apply the post-review attribution audit to the three authoritative DOCX files."""

from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
WEEK4 = ROOT / "SafeFAME-TS_01_选题与初步技术方案（第4周）.docx"
WEEK10 = ROOT / "SafeFAME-TS_02_书面中期进展报告（第10周）.docx"
FINAL = ROOT / "SafeFAME-TS_03_课程设计最终报告.docx"
SUMMARY = json.loads((ROOT / "outputs" / "attribution_audit_v4" / "summary.json").read_text(encoding="utf-8"))
AUDIT_FIGURE = ROOT / "outputs" / "attribution_audit_v4" / "fig_selected_route_attribution.png"
LEGACY_FIGURE = ROOT / "outputs" / "reviewer_sensitivity" / "fig_numeric_residual_ablation.png"


REPLACEMENTS = {
    "匹配容量纯数值残差对照": "纯数值残差消融对照",
    "匹配容量数值残差对照": "纯数值残差消融对照",
    "匹配数值残差对照": "纯数值残差消融对照",
    "匹配容量对照": "纯数值残差消融",
    "相对匹配容量纯数值残差对照": "相对纯数值残差消融对照",
    "文本增量归因检查": "文本相关增量诊断",
    "匹配对照区间正 / 负": "纯数值残差消融区间正 / 负",
    "匹配区间仍为": "纯数值残差消融区间仍为",
    "相对匹配数值对照": "相对纯数值残差消融对照",
    "相对匹配对照": "相对纯数值残差消融对照",
    "在匹配对照下": "在纯数值残差消融对照下",
    "匹配对照仅约束模型家族和调参范围": "纯数值残差消融仅保持模型家族和调参范围一致",
    "DOCX PDF 代码 数据清单 校验报告": "DOCX 代码 数据清单 校验报告；PDF定稿后生成",
    "DOCX PDF 代码 数据清单": "DOCX 代码 数据清单；PDF定稿后生成",
    "最终课程设计报告 DOCX 与 PDF；": "最终课程设计报告 DOCX；PDF在定稿并收到明确指令后生成；",
}

EXACT_REPLACEMENTS = {
    "以下补记形成于扩展v4完成后的代码—结果审计，不冒充第10周已完成内容。审计保留120项主实验和10次冻结路由，专门检验原对照是否足以把候选收益归因于文本。":
        "以下补记形成于扩展v4后的代码—结果审计，不冒充第10周已完成内容。表5汇总四组对照；审计保留120项主实验和10次冻结路由，检验原对照能否支持文本归因。",
    "表5汇总四组归因对照；原`matched_*`结果现统一解释为纯数值残差消融，而非匹配容量对照。等维对照只匹配设计宽度，不保证相同有效秩。唯一频率启用路线重新拟合置换交互尺度后，逐行p由0.021升至0.085；240条p值事后Holm校正无一达到0.05，BH校正有22条达到0.05。以上均不回写冻结门控。":
        "原`matched_*`结果现统一解释为纯数值残差消融，而非匹配容量对照。等维对照只匹配设计宽度，不保证相同有效秩。唯一频率启用路线重新拟合置换交互尺度后，逐行p由0.021升至0.085；240条p值事后Holm校正无一达到0.05，BH校正有22条达到0.05。以上均不回写冻结门控。",
    "表5汇总四组归因对照；原`matched_*`结果现统一解释为纯数值残差消融，而非纯数值残差消融。等维对照只匹配设计宽度，不保证相同有效秩。唯一频率启用路线重新拟合置换交互尺度后，逐行p由0.021升至0.085；240条p值事后Holm校正无一达到0.05，BH校正有22条达到0.05。以上均不回写冻结门控。":
        "原`matched_*`结果现统一解释为纯数值残差消融；等维对照只匹配设计宽度，不保证相同有效秩。唯一频率启用路线重新拟合置换交互尺度后，逐行p由0.021升至0.085；240条p值事后Holm校正无一达到0.05，BH校正有22条达到0.05。以上均不回写冻结门控。",
}


def set_run_font(run, chinese: str = "宋体", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), chinese)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def style_paragraph(paragraph, heading: bool = False, size: float | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, "黑体" if heading else "宋体", size=size, bold=True if heading else None)


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_terms(document: Document) -> None:
    for paragraph in all_paragraphs(document):
        if paragraph.text in EXACT_REPLACEMENTS:
            replacement = EXACT_REPLACEMENTS[paragraph.text]
            paragraph.clear()
            paragraph.add_run(replacement)
            style_paragraph(paragraph, paragraph.style.name.startswith("Heading"))
            continue
        if not any(old in paragraph.text for old in REPLACEMENTS):
            continue
        changed = False
        for run in paragraph.runs:
            for old, new in REPLACEMENTS.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    changed = True
        if any(old in paragraph.text for old in REPLACEMENTS):
            text = paragraph.text
            for old, new in REPLACEMENTS.items():
                text = text.replace(old, new)
            paragraph.clear()
            paragraph.add_run(text)
            changed = True
        if changed:
            style_paragraph(paragraph, paragraph.style.name.startswith("Heading"))


def find_paragraph(document: Document, exact: str | None = None, starts: str | None = None):
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if exact is not None and value == exact:
            return paragraph
        if starts is not None and value.startswith(starts):
            return paragraph
    raise ValueError(exact or starts)


def remove_existing_section(document: Document, heading_text: str, stop_text: str) -> None:
    heading = next((p for p in document.paragraphs if p.text.strip() == heading_text), None)
    if heading is None:
        return
    node = heading._p
    while node is not None:
        following = node.getnext()
        text = "".join(node.itertext()).strip()
        if node is not heading._p and text.startswith(stop_text):
            break
        node.getparent().remove(node)
        node = following


def insert_paragraph(target, text: str, style: str = "Body Text", bold_lead: str | None = None):
    paragraph = target.insert_paragraph_before(style=style)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, "宋体", bold=True)
        remainder = paragraph.add_run(text[len(bold_lead):])
        set_run_font(remainder)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, "黑体" if style.startswith("Heading") else "宋体",
                     bold=True if style.startswith("Heading") else None)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge_name, edge in edges.items():
        tag = f"w:{edge_name}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edge.items():
            element.set(qn(f"w:{key}"), str(value))


def insert_three_line_table(document: Document, target, headers: list[str], rows: list[list[str]],
                            font_size: float = 8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    table.autofit = True
    for col, text in enumerate(headers):
        table.rows[0].cells[col].text = text
    for values in rows:
        cells = table.add_row().cells
        for col, text in enumerate(values):
            cells[col].text = text
    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            tr_pr = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            for run in paragraph.runs:
                set_run_font(run, "黑体" if row_index == 0 else "宋体", size=font_size,
                             bold=row_index == 0)
            edges = {}
            if row_index == 0:
                edges["top"] = {"val": "single", "sz": "12", "color": "000000"}
                edges["bottom"] = {"val": "single", "sz": "8", "color": "000000"}
            if row_index == len(table.rows) - 1:
                edges["bottom"] = {"val": "single", "sz": "12", "color": "000000"}
            set_cell_border(cell, **edges)
    target._p.addprevious(table._tbl)
    return table


def apply_accessibility_fixes(document: Document, alt_by_id: dict[str, str] | None = None) -> None:
    for table in document.tables:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
    for doc_pr in document.element.body.xpath(".//wp:docPr"):
        description = (alt_by_id or {}).get(doc_pr.get("id"))
        if description:
            doc_pr.set("descr", description)


def insert_picture(document: Document, target, path: Path, width: float = 6.2):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    target._p.addprevious(paragraph._p)
    return paragraph


def replace_preceding_picture(document: Document, caption_starts: str, path: Path) -> None:
    caption = find_paragraph(document, starts=caption_starts)
    previous = caption._p.getprevious()
    if previous is not None and previous.xpath(".//w:drawing"):
        previous.getparent().remove(previous)
    insert_picture(document, caption, path)


def mark_fields_dirty(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    for field in document.element.body.xpath(".//w:fldChar[@w:fldCharType='begin']"):
        field.set(qn("w:dirty"), "true")


def update_week4(path: Path) -> None:
    document = Document(path)
    replace_terms(document)
    remove_existing_section(document, "审查后对照实验补记", "参考文献")
    target = find_paragraph(document, exact="参考文献")
    insert_paragraph(target, "审查后对照实验补记", "Heading 2")
    insert_paragraph(target, "以下内容为扩展v4完成后的审查后补记，不属于第4周预注册方案，也不回写当时的研究进度。原方案中的纯数值残差对照只保持Ridge模型家族和正则网格一致，输入维数并不相同，因此统一改称“纯数值残差消融”，不再称为匹配容量对照。")
    insert_paragraph(target, "对冻结门控实际启用的10条路线新增回退锚定因子对照：分别保留纯数值、文本质量特征、文本语义、完整特征，并设置与完整候选输入宽度相同的纯数值随机傅里叶对照。结果显示完整特征相对等维纯数值对照有9/10条点改善、6条区间为正；但语义单独相对纯数值只有4条区间为正，质量特征单独也有4条区间为正，故联合错位通过不能全部解释为自然语言语义贡献。")
    insert_paragraph(target, "唯一启用的频率路线在每次置换重新拟合语义—频率交互尺度后，逐行经验p由0.021变为0.085，不再达到0.025。冻结路由和原测试结果保留，但该路线的文本特异性解释降级。新增分析均为事后诊断，不改变99/999/5000次实验各自的原定角色。")
    mark_fields_dirty(document)
    document.save(path)


def update_week10(path: Path) -> None:
    document = Document(path)
    replace_terms(document)
    remove_existing_section(document, "审查后对照实验与结果补记", "参考文献")
    target = find_paragraph(document, exact="参考文献")
    insert_paragraph(target, "审查后对照实验与结果补记", "Heading 2")
    insert_paragraph(target, "以下补记形成于扩展v4后的代码—结果审计，不冒充第10周已完成内容。表5汇总四组对照；审计保留120项主实验和10次冻结路由，检验原对照能否支持文本归因。")
    insert_paragraph(target, "表 5 审查后启用路线归因对照", "Caption")
    insert_three_line_table(document, target,
        ["对照问题", "点改善", "区间正", "区间负"],
        [["回退锚定完整特征 对 实际回退", "6/10", "6/10", "1/10"],
         ["完整特征 对 等维纯数值", "9/10", "6/10", "0/10"],
         ["语义特征 对 纯数值", "7/10", "4/10", "0/10"],
         ["质量特征 对 纯数值", "9/10", "4/10", "0/10"]], font_size=9)
    insert_paragraph(target, "原`matched_*`结果现统一解释为纯数值残差消融；等维对照只匹配设计宽度，不保证相同有效秩。唯一频率启用路线重新拟合置换交互尺度后，逐行p由0.021升至0.085；240条p值事后Holm校正无一达到0.05，BH校正有22条达到0.05。以上均不回写冻结门控。")
    insert_paragraph(target, "若事后同时要求测试点改善、冻结区间为正、循环移位通过、回退锚定区间为正且等维对照区间为正，仅Agriculture H6 F1语义路线满足。该联合条件是在查看结果后定义，只用于说明证据强弱，不能称为预注册判据或因果确认。")
    mark_fields_dirty(document)
    document.save(path)


def update_final(path: Path) -> None:
    document = Document(path)
    replace_terms(document)
    remove_existing_section(document, "9.6 审查后对照实验与归因边界", "10 模型评价与结论")
    target = find_paragraph(document, exact="10 模型评价与结论")
    insert_paragraph(target, "9.6 审查后对照实验与归因边界", "Heading 2")
    insert_paragraph(target, "原v4输出中的`matched_*`对照只保持Ridge家族和正则网格一致。语义候选与纯数值对照的输入宽度分别为L+34和L，频率候选与对照分别为L+236和L+10，因此本报告统一将其改称纯数值残差消融。它可以检查完整特征包相对较简单数值残差模型是否有增量，但不能称为严格匹配容量，也不能识别文本因果效应。")
    insert_paragraph(target, "为去除Last锚点与实际回退专家不一致的混杂，只对冻结门控实际启用的10条路线建立二阶段回退锚定残差头。校准段拟合、决策段选择Ridge正则，再在校准加决策段重拟合并一次性评价测试段。五组输入分别为纯数值、数值加质量、数值加语义、完整特征和等维纯数值随机傅里叶特征。等维仅指输入宽度一致，不保证有效秩和归纳偏置完全相同。")
    insert_paragraph(target, "表 16 实际启用路线的审查后归因对照", "Caption")
    insert_three_line_table(document, target,
        ["对照问题", "点改善", "区间正", "区间负", "证据含义"],
        [["回退锚定完整特征 对 实际回退", "6/10", "6/10", "1/10", "去除Last锚点不对称"],
         ["完整特征 对 等维纯数值", "9/10", "6/10", "0/10", "匹配输入宽度"],
         ["语义特征 对 纯数值", "7/10", "4/10", "0/10", "不含质量特征"],
         ["质量特征 对 纯数值", "9/10", "4/10", "0/10", "文本元数据增量"],
         ["完整特征 对 质量特征", "7/10", "4/10", "0/10", "控制质量后的语义相关增量"],
         ["完整特征 对 语义特征", "9/10", "3/10", "0/10", "控制语义后的质量相关增量"]], font_size=7.8)
    insert_paragraph(target, "表16表明质量特征和语义特征都可能贡献收益，原逐行联合错位检验拒绝的是整个文本相关特征包的可交换性，不能把通过全部归因为自然语言语义。图16进一步对照冻结候选、回退锚定完整头和等维数值控制。")
    insert_picture(document, target, AUDIT_FIGURE, width=6.15)
    insert_paragraph(target, "图 16 实际启用路线的审查后归因诊断", "Caption")
    insert_paragraph(target, "频率候选的语义—频率交互尺度依赖文本与数值行的配对。原实现把对齐训练数据拟合的交互标准化器复用于置换数据；这不影响已保存计算的一致性，却削弱频率置换的严格交换性。对唯一启用的频率路线Agriculture H12 F1逐次重拟合交互尺度后，999次逐行p由0.021变为0.085，循环p由0.060变为0.068。冻结路由不追溯修改，但52.95%的测试点收益不再具有稳健的文本特异性支持。")
    insert_paragraph(target, "240条主门控p值事后整体校正时，原始p≤0.025有38条，Holm FWER 0.05下为0条，BH FDR 0.05下为22条；两种整体校正与逐任务两候选Bonferroni控制的错误范围不同，均不回写原门控。跨折描述中仅Climate H24和Environment H1在至少两折选择同一路线且至少两折点改善，滚动折共享历史，仍不构成独立重复。")
    insert_paragraph(target, "若事后同时要求原测试点改善、冻结选中路线区间为正、循环移位p≤0.025、回退锚定区间为正和等维对照区间为正，仅Agriculture H6 F1语义路线满足。该联合条件是在查看结果后定义的证据摘要，不是预注册门槛。总体结论因此收敛为：少数路线存在与文本相关的增量迹象，但大部分启用结果不能同时通过所有稳健性诊断。")

    conclusion_target = find_paragraph(document, exact="10.2 优点")
    insert_paragraph(conclusion_target, "审查后归因诊断没有改变120项冻结路由，但缩窄了可解释范围：10条启用路线中只有1条同时满足测试点改善、冻结区间、循环移位、回退锚定和等维数值对照五项事后条件。唯一频率启用路线的修正逐行p为0.085，因此其大幅点收益不能作为稳健语义证据。")
    replace_preceding_picture(document, "图 10 文本候选相对纯数值残差消融对照", LEGACY_FIGURE)
    mark_fields_dirty(document)
    document.save(path)


def main() -> None:
    if not AUDIT_FIGURE.is_file():
        raise FileNotFoundError(AUDIT_FIGURE)
    backup = ROOT / "paper" / "archive" / "20260920_before_attribution_audit"
    backup.mkdir(parents=True, exist_ok=True)
    for source in (WEEK4, WEEK10, FINAL):
        target = backup / source.name
        if not target.exists():
            shutil.copy2(source, target)
    if "--wording-only" in sys.argv:
        mirror = ROOT / "paper" / "final"
        mirror.mkdir(parents=True, exist_ok=True)
        for source in (WEEK4, WEEK10, FINAL):
            document = Document(source)
            replace_terms(document)
            alt_by_id = ({
                "16": "图10 文本候选相对纯数值残差消融对照的测试期差异",
                "17": "图16 实际启用路线的审查后归因诊断",
            } if source == FINAL else None)
            apply_accessibility_fixes(document, alt_by_id)
            mark_fields_dirty(document)
            document.save(source)
            shutil.copy2(source, mirror / source.name)
        print("wording updated:", WEEK4.name, WEEK10.name, FINAL.name)
        return
    if not LEGACY_FIGURE.is_file():
        raise FileNotFoundError(LEGACY_FIGURE)
    update_week4(WEEK4)
    update_week10(WEEK10)
    update_final(FINAL)
    mirror = ROOT / "paper" / "final"
    mirror.mkdir(parents=True, exist_ok=True)
    for source in (WEEK4, WEEK10, FINAL):
        shutil.copy2(source, mirror / source.name)
    print("updated:", WEEK4.name, WEEK10.name, FINAL.name)


if __name__ == "__main__":
    main()
