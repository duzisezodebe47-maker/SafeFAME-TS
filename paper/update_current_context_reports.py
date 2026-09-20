"""Strengthen current literature positioning and figure/table interpretation in the three reports."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORTS = [
    ROOT / "SafeFAME-TS_01_选题与初步技术方案（第4周）.docx",
    ROOT / "SafeFAME-TS_02_书面中期进展报告（第10周）.docx",
    ROOT / "SafeFAME-TS_03_课程设计最终报告.docx",
]

WEEK_REFS = [
    "[16] Ding K, Fan F, Wang Y, et al. DualSG: A Dual-Stream Explicit Semantic-Guided Multivariate Time Series Forecasting Framework. Proceedings of the 33rd ACM International Conference on Multimedia, 2025:508-517. DOI: 10.1145/3746027.3755458.",
    "[17] Liu P, Guo H, Dai T, et al. CALF: Aligning LLMs for Time Series Forecasting via Cross-modal Fine-Tuning. Proceedings of the AAAI Conference on Artificial Intelligence, 2025, 39(18):18915-18923. DOI: 10.1609/aaai.v39i18.34082.",
    "[18] Lin J, Wang Y, Luo H, Wang J, Pei Z. TiMi: Empower Time Series Transformers with Multimodal Mixture of Experts. arXiv preprint arXiv:2602.21693, 2026.",
    "[19] Jiang Y, Yu W, Lee G, et al. TimeXL: Explainable Multi-modal Time Series Prediction with LLM-in-the-Loop. Advances in Neural Information Processing Systems, 2025, 38.",
    "[20] Wu X, Jin J, Qiu W, et al. Aurora: Towards Universal Generative Multimodal Time Series Forecasting. International Conference on Learning Representations, 2026.",
    "[21] Tan M, Merrill M A, Gupta V, Althoff T, Hartvigsen T. Are Language Models Actually Useful for Time Series Forecasting? Advances in Neural Information Processing Systems, 2024, 37.",
]

FINAL_REFS = [
    "[19] Ding K, Fan F, Wang Y, et al. DualSG: A Dual-Stream Explicit Semantic-Guided Multivariate Time Series Forecasting Framework. Proceedings of the 33rd ACM International Conference on Multimedia, 2025:508-517. DOI: 10.1145/3746027.3755458.",
    "[20] Liu P, Guo H, Dai T, et al. CALF: Aligning LLMs for Time Series Forecasting via Cross-modal Fine-Tuning. Proceedings of the AAAI Conference on Artificial Intelligence, 2025, 39(18):18915-18923. DOI: 10.1609/aaai.v39i18.34082.",
    "[21] Tan M, Merrill M A, Gupta V, Althoff T, Hartvigsen T. Are Language Models Actually Useful for Time Series Forecasting? Advances in Neural Information Processing Systems, 2024, 37.",
]


def set_run_font(run, heading: bool = False, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "黑体" if heading else "宋体")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def style_paragraph(paragraph, heading: bool = False, size: float | None = None) -> None:
    for run in paragraph.runs:
        set_run_font(run, heading=heading, size=size, bold=True if heading else None)


def find_paragraph(document: Document, exact: str | None = None, starts: str | None = None):
    for paragraph in document.paragraphs:
        value = paragraph.text.strip()
        if exact is not None and value == exact:
            return paragraph
        if starts is not None and value.startswith(starts):
            return paragraph
    raise ValueError(exact or starts)


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)
    style_paragraph(paragraph, heading=paragraph.style.name.startswith("Heading"))


def insert_before(target, text: str, style: str = "Body Text"):
    paragraph = target.insert_paragraph_before(style=style)
    run = paragraph.add_run(text)
    is_heading = style.startswith("Heading")
    set_run_font(run, heading=is_heading, bold=True if is_heading else None)
    paragraph.paragraph_format.space_after = Pt(6)
    return paragraph


def body_text(document: Document) -> str:
    values = [p.text for p in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(values)


def set_cell_text(cell, text: str, header: bool = False) -> None:
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        for run in paragraph.runs:
            set_run_font(run, heading=header, size=8.2, bold=header)


def table_by_first_header(document: Document, header: str):
    matches = [table for table in document.tables if table.rows and table.rows[0].cells[0].text.strip() == header]
    if len(matches) != 1:
        raise ValueError(f"Expected one table headed {header!r}, found {len(matches)}")
    return matches[0]


def add_table_row(table, values: list[str]) -> None:
    cells = table.add_row().cells
    if len(cells) != len(values):
        raise ValueError("Table width mismatch")
    for cell, value in zip(cells, values, strict=True):
        set_cell_text(cell, value)


def update_table_row(table, key: str, values: list[str]) -> None:
    rows = [row for row in table.rows[1:] if row.cells[0].text.strip() == key]
    if len(rows) != 1:
        raise ValueError(f"Expected one row keyed by {key!r}, found {len(rows)}")
    if len(rows[0].cells) != len(values):
        raise ValueError("Table width mismatch")
    for cell, value in zip(rows[0].cells, values, strict=True):
        set_cell_text(cell, value)


def append_references(document: Document, references: list[str], before: str | None = None) -> None:
    existing = body_text(document)
    reference_paragraphs = [p for p in document.paragraphs if p.text.strip().startswith("[")]
    if not reference_paragraphs:
        raise ValueError("No existing reference paragraph found")
    style = reference_paragraphs[-1].style.name
    anchor = find_paragraph(document, exact=before) if before else None
    for text in references:
        if text.split(".", 1)[0] in existing:
            continue
        paragraph = anchor.insert_paragraph_before(style=style) if anchor is not None else document.add_paragraph(style=style)
        paragraph.add_run(text)
        style_paragraph(paragraph)


def mark_fields_dirty(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    for field in document.element.body.xpath(".//w:fldChar[@w:fldCharType='begin']"):
        field.set(qn("w:dirty"), "true")


def update_week4(path: Path) -> None:
    document = Document(path)
    if "近期多模态时序研究" in body_text(document):
        raise RuntimeError("Week 4 report already contains this revision")

    intro = find_paragraph(document, starts="现实预测通常同时面对")
    replace_paragraph(
        intro,
        "现实预测通常同时面对历史数值序列与新闻、报告、政策或事件文本。文本可能提前暴露供需、风险或制度变化，也可能晚于预测时点、偏离主题或只是复述已发生趋势。能源调度、交通管理、气候与农业预警等场景既需要利用外部信息，又不能接受未经验证的文本信号直接覆盖稳定数值模型。Time-MMD提供多领域成对数据，为这一问题提供实验基础[1]。本项目因此把文本视为受信息可得性约束的候选外生变量，而不是默认有效的额外模态。",
    )

    table = table_by_first_header(document, "研究")
    update_table_row(
        table,
        "Time-LLM / TimeCMA",
        ["Time-LLM / TimeCMA / CALF", "ICLR 2024 / AAAI 2025", "重编程与跨模态分布对齐", "说明融合方法空间；不复现LLM微调"],
    )
    add_table_row(table, ["DualSG", "ACM MM 2025", "数值预测流与语义指导流协作", "支持残差修正定位；不复现LLM推理流"])
    add_table_row(table, ["TiMi / TimeXL / Aurora", "2026预印本 / NeurIPS 2025 / ICLR 2026", "动态路由 解释闭环 跨域基础模型", "作为研究前沿对照；不写成已实现模块"])
    add_table_row(table, ["Tan等", "NeurIPS 2024", "检验语言模型是否提供独立时序增益", "支持错位 容量与数值消融对照"])

    related = find_paragraph(document, starts="数值侧从 Last")
    replace_paragraph(
        related,
        "近期多模态时序研究从语言重编程发展到跨模态对齐、显式双流、时频文本融合、混合专家路由和大模型在环解释。数值模型背景来自PatchTST、DLinear、TimesNet、Crossformer、FEDformer、iTransformer与TimeXer[2-8]，语义编码和重编程背景来自Time-LLM、Sentence-BERT与MiniLM[9-11]。TimeCMA与CALF关注时序和语言表征的错配[12,17]；DualSG把语言模型放在语义指导位置，而不是替代数值预测器[16]；T3Time与Spectral Text Fusion强调跨度条件和频率尺度[13-14]；TiMi、TimeXL与Aurora分别代表动态路由、解释闭环和跨域预训练[18-20]。与此同时，Tan等的研究提示，语言模型带来的提升必须与更简单模型和容量因素区分[21]。表1据此说明本项目的取舍。",
    )
    heading3 = find_paragraph(document, exact="3 数据与任务设计")
    insert_before(
        heading3,
        "本项目不把上述大型结构拼接成新网络，而把研究问题收缩为一个可证伪的部署判断：在严格时点条件下，文本相关特征是否在验证期稳定胜过数值回退和错位对照。冻结MiniLM、线性残差头、分段门槛与安全回退共同降低容量混杂；频率交互只检验多尺度文本影响的轻量版本，不声称复现T3Time或SpecTF。",
    )
    heading4 = find_paragraph(document, exact="4 初步算法方案")
    insert_before(
        heading4,
        "表3把主实验、外部数值验证和扰动测试分开：Time-MMD回答文本是否提供增量，ETT只核查数值训练管线，输入扰动只诊断数值模型的失效方式。三类证据回答的问题不同，不能相互替代。",
    )
    append_references(document, WEEK_REFS)
    document.core_properties.comments = "Current literature positioning and figure/table interpretation revised; experimental results unchanged"
    mark_fields_dirty(document)
    document.save(path)


def update_week10(path: Path) -> None:
    document = Document(path)
    if "1.1 近期研究对照与阶段取舍" in body_text(document):
        raise RuntimeError("Week 10 report already contains this revision")

    stage = find_paragraph(document, exact="阶段完成度")
    insert_before(
        stage,
        "图1从左到右对应实际执行顺序：数值和文本先按预测时点形成样本，数值专家与两条文本残差候选分别训练，验证期门控只在满足总体、分段和错位检验条件时启用文本，否则输出数值回退。该流程把文本融合转化为可审计的路径选择，而不是默认融合。",
    )
    insert_before(stage, "1.1 近期研究对照与阶段取舍", "Heading 2")
    insert_before(
        stage,
        "截至第10周，多模态时序研究已出现三类更强结构：TimeCMA与CALF进行跨模态对齐[12,17]，DualSG以数值流承担精确预测并让语义流做修正[16]，T3Time与SpecTF在时域和频域融合文本[13-14]。TiMi、TimeXL和Aurora进一步引入混合专家路由、解释闭环和跨域预训练[18-20]。这些工作说明文本建模正在从简单拼接转向条件化使用，但也增加了模型容量和归因难度。",
    )
    insert_before(
        stage,
        "本项目在课程设计规模内选择轻量而可证伪的路线：冻结文本编码器，用Ridge残差头隔离增量，并要求对齐文本胜过错位文本后才允许启用。Tan等对语言模型时序价值的审慎评估[21]也支持保留强数值基线和反事实对照。现实上，这种门控适合能源、交通、农业和安全等误启用代价较高的预测流程：文本证据不足时系统仍能给出已有数值模型结果，并保留启用理由供复核。",
    )
    heading2 = find_paragraph(document, exact="2 已完成的数据工作")
    insert_before(
        heading2,
        "表1显示第10周已经完成数据、文本索引、数值基线和候选实现，真正冻结的是选择协议；报告与原型仍在整理。这里的“完成”指已有可检查文件和自检记录，不代表最终测试、扩展滚动回测或审查后归因分析在第10周已经完成。",
    )
    append_references(document, WEEK_REFS)
    document.core_properties.comments = "Current literature positioning and figure/table interpretation revised; experimental results unchanged"
    mark_fields_dirty(document)
    document.save(path)


def update_final(path: Path) -> None:
    document = Document(path)
    if "SafeFAME-TS的创新位置" in body_text(document):
        raise RuntimeError("Final report already contains this revision")

    heading12 = find_paragraph(document, exact="1.2 闭环对应关系")
    insert_before(
        heading12,
        "现实应用中，能源负荷可能受政策和天气通报影响，交通流可能受事故和管制信息影响，农业与公共健康预测也会受到外部事件驱动。文本若及时且相关，可能补充数值历史尚未呈现的变化；若过期、重复或主题漂移，则可能造成负迁移。本文关注的现实意义因此是建立一个可审计的启用条件：文本证据不足时保留稳定数值预测，证据较强时才允许文本相关残差修正。该机制提供的是风险受控的模型使用方式，不是对具体事件的因果解释。",
    )
    heading2 = find_paragraph(document, exact="2 相关研究与方案依据")
    insert_before(
        heading2,
        "表1把课程要求落实到数据、模型、输出和验证证据。18项历史主实验负责建立完整闭环，后续v3纠错和v4扩展则检验同一逻辑在更严格边界和更多滚动任务下是否仍成立；因此正文同时保留历史结果与审查后结果，不用后者改写前者的时间顺序。",
    )

    table = table_by_first_header(document, "研究")
    update_table_row(
        table,
        "Time-LLM / TimeCMA",
        ["Time-LLM / TimeCMA / CALF", "ICLR 2024 / AAAI 2025", "重编程与跨模态分布对齐", "说明融合空间；不复现LLM微调"],
    )
    update_table_row(
        table,
        "TimeXL",
        ["TimeXL", "NeurIPS 2025", "原型解释与大模型反思闭环", "启发案例审计；不复制三LLM结构"],
    )
    add_table_row(table, ["DualSG", "ACM MM 2025", "数值流与语义指导流协作", "支持残差修正定位；不复现LLM推理流"])
    add_table_row(table, ["Tan等", "NeurIPS 2024", "检验语言模型是否提供独立时序增益", "支持错位 容量与数值消融对照"])

    related1 = find_paragraph(document, starts="PatchTST 的分块")
    replace_paragraph(
        related1,
        "数值侧由PatchTST的分块和通道独立结构[2]、DLinear的强简单基线[3]、TimesNet与FEDformer的多周期或频域建模[4,6]构成；Crossformer、iTransformer与TimeXer分别提供跨维依赖、变量中心表示和外生变量背景[5,7-8]。文本侧用Sentence-BERT和MiniLM提供冻结句向量[10-11]。Time-LLM、TimeCMA与CALF展示了重编程、跨模态相似度对齐和多层分布对齐的不同路线[9,12,20]。这些工作追求更强表示，本项目则用冻结编码和线性残差头减少模型容量对“文本增量”判断的干扰。（见表2）",
    )
    related2 = find_paragraph(document, starts="相关多模态工作包括")
    replace_paragraph(
        related2,
        "近期研究进一步转向条件化使用文本。DualSG让数值流承担精细预测、语义流提供趋势修正[19]；T3Time和SpecTF分别用跨度条件门控与频域交互处理多尺度影响[13-14]；TiMi用混合专家选择性路由文本指导[16]；TimeXL强调预测、反思和文本修订形成的解释闭环[17]；Aurora把文本和图像纳入跨域预训练与生成式预测[18]。Tan等的反证式研究则提醒，语言模型改进必须排除更简单模型和容量差异[21]。这些结果共同说明，当前问题已从“能否融合”推进到“何时融合、如何证明增量、失败时如何回退”。",
    )
    related3 = find_paragraph(document, starts="TimeCMA 已正式发表于")
    replace_paragraph(
        related3,
        "SafeFAME-TS的创新位置是验证与路由框架，而不是新的大规模预测骨干。它实际借鉴了四类思想：将文本作为外生条件，保留时域和频域两类候选，按任务与跨度做条件选择，并为选择结果保留解释记录。它没有复现CALF的LLM微调、DualSG的生成式语义推理、TiMi的神经混合专家、TimeXL的三LLM循环或Aurora的基础模型预训练。由此，本文的结论只约束当前轻量表示和冻结门控，不能否定这些大型模型在其他数据和训练协议下的效果。",
    )

    version = find_paragraph(document, starts="版本说明：原始v2")
    insert_before(
        version,
        "表4集中区分观测量、构造特征、模型参数和决策统计量。尤其是Zt与qt分别表示语义内容和文本质量，b与r分别表示数值回退和文本候选；后续消融因此能够检查语义、质量和数值残差结构的不同贡献。",
    )
    heading51 = find_paragraph(document, exact="5.1 数值专家池")
    insert_before(
        heading51,
        "图2中的门控位于候选训练之后、测试评估之前。它只决定本任务使用文本候选还是数值回退，不在单个预测步上动态混合，也不根据测试期表现反向修改阈值；这一区别使冻结选择可以从保存的校准、决策和测试证据中复核。",
    )
    heading6 = find_paragraph(document, exact="6 实验设计与求解")
    insert_before(
        heading6,
        "表5给出从原始文件到冻结评估的完整顺序。步骤3至5只使用训练、校准和路径决策数据形成回退与资格标记；步骤6在冻结参数下重拟合；步骤7才读取测试目标。审查后v3与v4继续沿用这一用途分离，但修复目标窗口跨界并扩大置换次数。",
    )
    heading63 = find_paragraph(document, exact="6.3 评价指标")
    insert_before(
        heading63,
        "表6表明深度模型、Ridge、置换和区间估计使用不同的随机或重复设置。三个深度种子用于历史v2模型平均，v4扩展为五个；99或999次错位决定主门控的经验p值分辨率，5000次Bootstrap只刻画测试损失差的不确定性，重复次数增加不等于增加独立样本。",
    )
    heading14 = find_paragraph(document, starts="表 14 v4候选")
    insert_before(
        heading14,
        "图13中每一折都按训练、校准、决策和测试的目标区间顺序推进，后折可以使用此前已观测历史，但三折测试目标区间不重叠。该设计增加时间位置覆盖，却不能把共享历史的折视为独立数据集。",
    )
    after15 = find_paragraph(document, starts="此次扩展是在已有历史数据上")
    insert_before(
        after15,
        "图15显示大多数任务—折达到8个完整非重叠决策块的参考线，但低频长跨度任务的有效块数仍远少于窗口数。图中的参考线用于暴露功效差异，不是新的放行条件，也不改变任何冻结路由。",
    )

    old_ref = find_paragraph(document, starts="[17] Jiang Y")
    replace_paragraph(old_ref, "[17] Jiang Y, Yu W, Lee G, et al. TimeXL: Explainable Multi-modal Time Series Prediction with LLM-in-the-Loop. Advances in Neural Information Processing Systems, 2025, 38.")
    append_references(document, FINAL_REFS, before="附录 A 关键结果复核")
    document.core_properties.comments = "Current literature positioning and figure/table interpretation revised; experimental results unchanged"
    mark_fields_dirty(document)
    document.save(path)


def main() -> None:
    update_week4(REPORTS[0])
    update_week10(REPORTS[1])
    update_final(REPORTS[2])
    mirror = ROOT / "paper" / "final"
    mirror.mkdir(parents=True, exist_ok=True)
    for report in REPORTS:
        shutil.copy2(report, mirror / report.name)
    print("Updated three reports and synchronized paper/final mirrors")


if __name__ == "__main__":
    main()
