"""Complete in-text citation coverage after the current-literature revision."""

from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]


def set_font(paragraph) -> None:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "宋体")


def replace_once(path: Path, old: str, new: str) -> None:
    document = Document(path)
    matches = [p for p in document.paragraphs if p.text == old]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph in {path.name}, found {len(matches)}")
    paragraph = matches[0]
    paragraph.clear()
    paragraph.add_run(new)
    set_font(paragraph)
    document.save(path)


def main() -> None:
    week4 = ROOT / "SafeFAME-TS_01_选题与初步技术方案（第4周）.docx"
    final = ROOT / "SafeFAME-TS_03_课程设计最终报告.docx"
    old_week4 = "近期多模态时序研究从语言重编程发展到跨模态对齐、显式双流、时频文本融合、混合专家路由和大模型在环解释。TimeCMA与CALF关注时序和语言表征的错配[12,17]；DualSG把语言模型放在语义指导位置，而不是替代数值预测器[16]；T3Time与Spectral Text Fusion强调跨度条件和频率尺度[13-14]；TiMi、TimeXL与Aurora分别代表动态路由、解释闭环和跨域预训练[18-20]。与此同时，Tan等的研究提示，语言模型带来的提升必须与更简单模型和容量因素区分[21]。表1据此说明本项目的取舍。"
    new_week4 = "近期多模态时序研究从语言重编程发展到跨模态对齐、显式双流、时频文本融合、混合专家路由和大模型在环解释。数值模型背景来自PatchTST、DLinear、TimesNet、Crossformer、FEDformer、iTransformer与TimeXer[2-8]，语义编码和重编程背景来自Time-LLM、Sentence-BERT与MiniLM[9-11]。TimeCMA与CALF关注时序和语言表征的错配[12,17]；DualSG把语言模型放在语义指导位置，而不是替代数值预测器[16]；T3Time与Spectral Text Fusion强调跨度条件和频率尺度[13-14]；TiMi、TimeXL与Aurora分别代表动态路由、解释闭环和跨域预训练[18-20]。与此同时，Tan等的研究提示，语言模型带来的提升必须与更简单模型和容量因素区分[21]。表1据此说明本项目的取舍。"
    old_final = "数值侧由PatchTST的分块和通道独立结构[2]、DLinear的强简单基线[3]、TimesNet与FEDformer的多周期或频域建模[4,6]构成；iTransformer与TimeXer分别提供变量中心表示和外生变量背景[7-8]。文本侧用Sentence-BERT和MiniLM提供冻结句向量[10-11]。Time-LLM、TimeCMA与CALF展示了重编程、跨模态相似度对齐和多层分布对齐的不同路线[9,12,20]。这些工作追求更强表示，本项目则用冻结编码和线性残差头减少模型容量对“文本增量”判断的干扰。（见表2）"
    new_final = "数值侧由PatchTST的分块和通道独立结构[2]、DLinear的强简单基线[3]、TimesNet与FEDformer的多周期或频域建模[4,6]构成；Crossformer、iTransformer与TimeXer分别提供跨维依赖、变量中心表示和外生变量背景[5,7-8]。文本侧用Sentence-BERT和MiniLM提供冻结句向量[10-11]。Time-LLM、TimeCMA与CALF展示了重编程、跨模态相似度对齐和多层分布对齐的不同路线[9,12,20]。这些工作追求更强表示，本项目则用冻结编码和线性残差头减少模型容量对“文本增量”判断的干扰。（见表2）"
    replace_once(week4, old_week4, new_week4)
    replace_once(final, old_final, new_final)
    mirror = ROOT / "paper" / "final"
    for report in ROOT.glob("SafeFAME-TS_0*.docx"):
        shutil.copy2(report, mirror / report.name)
    print("Citation coverage completed and mirrors synchronized")


if __name__ == "__main__":
    main()
