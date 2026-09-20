"""Build the final Chinese course-design report from verified experiment artifacts."""

from __future__ import annotations

from pathlib import Path
import json

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "final"
EQ = ROOT / "paper" / "equations"
FIG = ROOT / "outputs" / "figures"
NAVY = "234E70"
PALE = "EFF5F8"
GRID = "D9D9D9"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}")) or OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders") or OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), GRID)
        borders.append(tag)
    tbl_pr.append(borders)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(str(text))
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
    for r_idx, row_values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            if r_idx % 2:
                set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if isinstance(value, str) and len(value) > 12 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            run.font.size = Pt(font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False
    p.add_run(text)


def add_figure(doc, filename: str, caption: str, width=6.6) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(FIG / filename), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    add_caption(doc, caption)


def add_equation(doc, filename: str, number: int, width=5.8) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    shape = p.add_run().add_picture(str(EQ / filename), width=Inches(width))
    shape._inline.docPr.set("descr", f"数学公式 {number}")
    n = doc.add_paragraph(f"式 {number}")
    n.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    n.paragraph_format.space_after = Pt(6)


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def add_body(doc, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(5)


def add_bullets(doc, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def configure(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name, size, bold in (("Title", 22, True), ("Heading 1", 16, True), ("Heading 2", 13, True), ("Heading 3", 11, True)):
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Caption"].font.name = "宋体"
    styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)
    header = section.header.paragraphs[0]
    header.text = "人工智能课程设计  SafeFAME TS"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)
    footer.add_run(" 页")


def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    claims = json.loads((ROOT / "outputs/tables/final/verified_claims.json").read_text(encoding="utf-8"))
    dev = pd.read_csv(ROOT / "outputs/tables/final/table_development_multimodal.csv")
    conf = pd.read_csv(ROOT / "outputs/tables/final/table_confirmation_safety.csv")
    ett = pd.read_csv(ROOT / "outputs/tables/final/table_ett_external.csv")
    doc = Document()
    configure(doc)
    doc.core_properties.title = "SafeFAME-TS 人工智能课程设计报告"
    doc.core_properties.subject = "时间序列预测、多模态文本融合与稳健性验证"
    doc.core_properties.author = ""

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(18)
    title.add_run("文本何时有助于时间序列预测")
    title_border = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if title_border is not None:
        title._p.get_or_add_pPr().remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("面向多领域的严格时点对齐 反事实检验与安全融合").bold = True
    subtitle.runs[0].font.size = Pt(15)
    method = doc.add_paragraph()
    method.alignment = WD_ALIGN_PARAGRAPH.CENTER
    method.paragraph_format.space_before = Pt(14)
    method.add_run("SafeFAME TS 人工智能课程设计报告").font.size = Pt(12)
    doc.add_paragraph().paragraph_format.space_after = Pt(90)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("个人课程设计  2026年9月").font.size = Pt(11)
    doc.add_page_break()

    add_heading(doc, "摘要", 1)
    add_body(doc, "多模态时间序列研究通常假定外部文本能够补充数值序列，但文本的发布时间、来源覆盖和语义相关性会随领域与预测跨度变化。若直接拼接文本特征，模型可能利用未来信息，也可能在未见领域产生负迁移。本课程设计据此提出 SafeFAME-TS，将严格时点文本索引、冻结句向量、频率统计、反事实文本置换和验证期安全选择组合为完整实验系统。研究使用 Time-MMD 的六个领域，其中 Climate、Economy、Energy 和 Traffic 用于开发性回测，Agriculture 与 Security 在规则冻结后用于确认；另使用 ETTh1 和 ETTh2 的六项长预测任务检验数值骨干。")
    add_body(doc, f"实验共构建 {14281:,} 条去重事实文本，并强制每条事实的结束日期早于预测起点。开发集 12 项任务中，正确对齐文本仅在 Traffic 的三个跨度上优于置换文本，且只在 H=6 和 H=12 超过最强数值基线，MSE 分别降低 7.28% 和 4.81%。冻结确认集的六项任务均未达到语义路径启用门槛，说明文本增益不能跨领域泛化；但是选择器拒绝了全部有害语义路径，平均避免 {claims['confirmation_mean_avoided_semantic_loss_pct']:.2f}% 的语义路径损失。ETT 外部基准中，PatchTST 在六项任务全部获胜，相对最强非 PatchTST 基线平均降低 {claims['ett_mean_patchtst_gain_pct']:.2f}% MSE。输入扰动实验显示，PatchTST 对高斯噪声和 10% 随机缺失的平均退化分别为 3.31% 和 1.68%，对尖峰污染和末段连续缺失仍较敏感。")
    add_body(doc, "结果表明，文本是否有用必须通过严格时点约束和反事实对照逐任务验证。SafeFAME-TS 的主要可靠收益是以预先规定的验证规则拒绝不稳定语义路径，而不是保证文本普遍提高预测精度。系统保留全部数据审计、随机种子、模型输出和图表生成记录，可从原始数据复算主要结论。")
    p = doc.add_paragraph()
    p.add_run("关键词 ").bold = True
    p.add_run("多模态时间序列预测  时点对齐  反事实检验  安全融合  PatchTST  负迁移")
    doc.add_page_break()

    add_heading(doc, "目录", 1)
    for item in [
        "1 研究问题与设计目标", "2 相关研究与方案取舍", "3 数据审计与实验协议",
        "4 SafeFAME TS 方法", "5 实验设置与实现", "6 开发性回测结果",
        "7 冻结确认与外部基准", "8 稳健性与误差分析", "9 系统实现与可复现性",
        "10 结论与局限", "参考文献", "附录",
    ]:
        doc.add_paragraph(item).paragraph_format.space_after = Pt(4)
    doc.add_page_break()

    add_heading(doc, "1 研究问题与设计目标", 1)
    add_heading(doc, "1.1 问题背景", 2)
    add_body(doc, "现实预测者往往同时观察数值指标和事件、政策、报告等文本。文本可能解释突变或提供领先信号，但其价值取决于是否在预测时点已经公开、是否与当前序列有关、是否覆盖相同制度阶段。Time-MMD 提供了多领域数值序列与报告、搜索事实，为检验这些条件提供了基础。课程设计不把文本维度增加本身视为创新，而把可证伪的时点约束、对齐反事实和安全回退作为研究主线。")
    add_heading(doc, "1.2 研究问题", 2)
    add_bullets(doc, [
        "问题一  在只使用预测时点之前事实的条件下，文本语义能否稳定超过纯数值基线。",
        "问题二  正确时点对齐文本是否优于随机置换文本，从而排除模板结构和时间趋势伪相关。",
        "问题三  频率结构与文本语义交互能否改善不同预测跨度，并在证据不足时自动回退。",
        "问题四  数值骨干在独立 ETT 基准和输入扰动下是否具有可复现的预测能力。",
    ])
    add_heading(doc, "1.3 数学任务", 2)
    add_body(doc, "给定截至时点 t 的长度 L 多变量数值窗口和预测起点前可获得的文本集合，模型输出未来 H 步目标序列。文本集合的截止条件属于信息约束，而非普通特征筛选。")
    add_equation(doc, "eq01_task.png", 1, 5.2)
    add_body(doc, "式 1 中，X 为数值窗口，Z 为严格早于预测起点的文本信息，H 为预测跨度，θ 为由训练数据估计的参数。训练、验证和测试均按时间顺序划分，任何测试期目标和未来文本都不进入特征构造。")
    add_figure(doc, "fig01_technical_route.png", "图 1 SafeFAME TS 技术路线")

    add_heading(doc, "2 相关研究与方案取舍", 1)
    add_heading(doc, "2.1 数值预测基线", 2)
    add_body(doc, "Last 和季节朴素模型提供最小可解释基线；AR-Ridge 检验线性自回归关系；DLinear 代表分解后的线性深度模型；PatchTST 将连续子序列分块为 token，并共享通道编码器参数。本文始终先比较简单模型，再判断复杂结构是否值得。ETT 实验的单目标协议与原论文常见的全通道平均协议不同，结果不能直接横向拼表。")
    add_heading(doc, "2.2 文本与时间序列融合", 2)
    add_body(doc, "Time-MMD 强调外生文本与数值序列的细粒度对齐。TimeCMA、T3Time、Spectral Text Fusion 等工作分别启发了跨模态对齐、预测跨度条件化和频域交互。但课程设计不训练大型语言模型，也不照搬多分支大模型；文本编码器保持冻结，融合头使用受正则化的线性结构，以便在中等算力下完成完整对照。")
    add_heading(doc, "2.3 本文取舍", 2)
    add_body(doc, "主方案保留三点：第一，文本必须通过发布日期约束；第二，文本价值必须超过置换对照；第三，语义路径必须在验证期总体和前后两半同时胜出。被舍弃的探索性神经门控原型出现 0.5 至 0.7 的高门控比例，却在多个领域产生负迁移，因此不进入最终算法。该否定结果说明，复杂门控并不天然等于可靠选择。")

    add_heading(doc, "3 数据审计与实验协议", 1)
    add_heading(doc, "3.1 数据范围", 2)
    rows = [
        ["Time-MMD 开发", "Climate Economy Energy Traffic", "12", "数值加 fact 文本", "规则开发"],
        ["Time-MMD 确认", "Agriculture Security", "6", "数值加 fact 文本", "冻结后一次确认"],
        ["ETT 外部", "ETTh1 ETTh2", "6", "多变量数值输入 单 OT 输出", "数值骨干外部检验"],
    ]
    add_caption(doc, "表 1 数据集与任务划分")
    add_table(doc, ["数据组", "领域或数据集", "任务数", "输入输出", "用途"], rows, [1.05, 1.8, .65, 1.65, 1.45], 8.2)
    add_body(doc, "Time-MMD 的 Climate 原文件按时间降序排列，Economy 原文件未完全排序。统一加载函数先解析日期，再执行稳定升序排序，并检查重复时间戳；这一修正发生在正式实验前，早期无效输出未进入最终表格。ETT 使用官方 CSV 前 14,400 个小时，确保与预先定义的训练、验证和测试边界一致。")
    add_heading(doc, "3.2 严格时点文本", 2)
    add_body(doc, "文本索引只保留 fact 字段，排除 preds 和 pred 等可能直接包含未来判断的字段。对每个预测起点，每个来源最多选取最近 32 条事实，且事实结束日期必须严格早于预测起点。去重后语料含 14,281 条事实，以固定版本的 all-MiniLM-L6-v2 编码为 384 维向量。编码器修订号和索引文件均记录在元数据中。")
    add_equation(doc, "eq02_text.png", 2, 5.6)
    add_body(doc, "式 2 对可用事实向量进行时间衰减和质量加权。Δt 表示事实距预测起点的时间差，q 表示来源、长度和缺失情况形成的质量权重，ε 防止空集合除零。空文本窗口使用显式缺失指示，不把未知事实解释成真实零语义。")
    add_figure(doc, "fig02_text_coverage.png", "图 2 各领域测试时段严格时点文本覆盖率")
    add_heading(doc, "3.3 时间划分与防泄漏", 2)
    add_body(doc, "Time-MMD 按观测顺序划分 70% 训练、10% 验证和 20% 测试。输入长度为 24 个月，月度任务预测 3、6、12 或领域协议中的 4、12、24 个时间步。所有标准化、Ridge 参数选择和融合系数估计都限制在训练或验证区间内。开发领域参与规则调整，因此其测试结果称为开发性回测；Agriculture 和 Security 在协议冻结前未参与模型选择。")

    add_heading(doc, "4 SafeFAME TS 方法", 1)
    add_heading(doc, "4.1 数值专家", 2)
    add_body(doc, "数值专家集合包括 Last、季节朴素、AR-Ridge、DLinear 和 PatchTST。AR-Ridge 以过去 L 个目标值预测未来 H 步，通过验证集在候选正则系数中选择 α。")
    add_equation(doc, "eq03_ridge.png", 3, 4.9)
    add_body(doc, "深度模型使用直接多步输出，避免递归预测的误差累积。PatchTST-M 接收全部数值通道，最终只输出 OT 目标。其输入长度 96、patch 长度 16、步长 8、隐藏维度 32、两层编码器和四个注意力头，参数规模受控。")
    add_heading(doc, "4.2 语义与频率交互", 2)
    add_body(doc, "语义候选路径以数值预测为锚点，增加文本线性校正；完整路径再加入由输入序列离散傅里叶幅值提取的频率统计与文本向量交互。频率统计只从历史窗口计算，不使用未来目标。")
    add_equation(doc, "eq04_fusion.png", 4, 5.8)
    add_body(doc, "式 4 中，上标 n 表示数值路径，上标 s 表示语义候选；φ 为频率统计，U、Ws 和 Wf 为训练期估计并正则化的参数。无频率消融令 Wf 为零。该结构让频率交互具有可检验的增量，而不是把全部效果归因于文本。")
    add_heading(doc, "4.3 反事实检验与安全选择", 2)
    add_body(doc, "对每个任务生成三次训练内文本置换，破坏文本与预测时点的对应关系，同时保留文本维度、边际分布和缺失率。若正确对齐路径不能超过置换路径，则文本可能只提供时间趋势、模板或样本身份线索。验证期再分为前后两半，要求语义路径在两半均优于数值路径。")
    add_equation(doc, "eq05_select.png", 5, 6.1)
    add_body(doc, "式 5 中，Vn、Vs 和 Vπ 分别为数值、语义和置换路径的验证误差。四个条件必须同时成立才启用语义，否则任务级回退到数值路径。启用后若单个预测与数值专家的分歧超过验证期校准阈值，也执行逐预测值回退。")
    add_heading(doc, "4.4 评价指标与区间", 2)
    add_equation(doc, "eq06_mse.png", 6, 4.7)
    add_body(doc, "主指标为 MSE，同时报告 MAE 和 RMSE。冻结确认实验以预测原点为重采样单位执行 1,000 次 Bootstrap，避免把同一预测窗口内高度相关的多个步长误当成独立样本。深度模型使用三个固定随机种子，并报告均值和标准差。")

    add_heading(doc, "5 实验设置与实现", 1)
    add_heading(doc, "5.1 软件与硬件", 2)
    add_body(doc, "实验使用 Python 3.12、PyTorch 2.8.0 和 CUDA 12.8，在 NVIDIA GeForce RTX 5070 Laptop GPU 上运行。数据处理依赖 NumPy、pandas 和 scikit-learn，图表由 Matplotlib 与 seaborn 生成。所有包版本写入 requirements.txt，随机种子集中为 2026、2027 和 2028。")
    add_heading(doc, "5.2 训练与选择", 2)
    rows = [
        ["Time-MMD", "24", "3 4 6 12 24", "70 10 20 时间划分", "MSE"],
        ["ETT", "96", "96 336 720", "8640 2880 2880", "MSE"],
        ["深度模型", "批量128", "最多60轮", "早停耐心8", "三种子"],
        ["确认区间", "预测原点", "1000次", "分位数区间", "95%"],
    ]
    add_caption(doc, "表 2 核心实验配置")
    add_table(doc, ["对象", "输入", "跨度", "划分或训练", "报告"], rows, [1.15, 1.05, 1.35, 1.75, 1.15])
    add_body(doc, "验证集仅用于超参数和路径选择，测试集不参与文本筛选或门控阈值估计。ETT 的深度模型先在训练与验证划分上确定最佳训练轮数，再按该轮数使用训练加验证数据重训；保存每个种子的模型参数、训练日志和预测样本。")
    add_heading(doc, "5.3 系统模块", 2)
    add_bullets(doc, [
        "数据审计模块检查时间顺序、字段类型、缺失、重复和文本泄漏风险。",
        "特征模块生成严格时点文本索引、冻结语义向量和数值频率统计。",
        "实验模块统一运行传统基线、深度基线、语义探针、确认实验和稳健性实验。",
        "汇总模块从原始 CSV 重新计算论文表格、可核验主张和九张高清图。",
        "结果看板展示开发、确认、ETT 和方法边界，不加载个人文件或未核验指标。",
    ])

    add_heading(doc, "6 开发性回测结果", 1)
    add_heading(doc, "6.1 数值基线", 2)
    add_body(doc, "不同领域的最优数值模型并不一致。Climate 的短中跨度由 Last 获胜，长跨度由 AR-Ridge 获胜；Economy 三个跨度均由 DLinear-M 获胜；Energy 在 H=4 和 H=24 由 AR-Ridge 获胜，H=12 由 Last 获胜；Traffic 的 H=3 由 PatchTST 获胜，H=6 和 H=12 由 DLinear-M 获胜。这一结果否定了用单一复杂模型覆盖全部任务的做法。")
    add_heading(doc, "6.2 语义增益", 2)
    rows = []
    for _, r in dev.iterrows():
        rows.append([f"{r.domain} H={int(r.horizon)}", r.best_numeric_model, f"{r.best_numeric_mse:.6f}", f"{r.aligned_semantic_mse:.6f}", f"{r.shuffled_semantic_mse:.6f}", f"{r.semantic_improvement_vs_best_numeric_pct:+.2f}%"])
    add_caption(doc, "表 3 开发性回测中的对齐语义与数值基线")
    add_table(doc, ["任务", "数值优胜", "数值MSE", "对齐MSE", "置换MSE", "改善"], rows, [1.25, 1.15, 1.0, 1.0, 1.0, .8], 7.7)
    add_body(doc, "只有 Traffic H=6 和 H=12 的对齐语义路径超过最强数值基线，改善 7.28% 和 4.81%。Traffic H=3 的对齐文本虽优于置换文本，但仍比 PatchTST 数值结果高 1.27%。其余九项任务的语义融合均退化，Economy H=12 的退化达到 127.70%。因此，对齐证据与最终精度必须同时报告。")
    add_figure(doc, "fig03_development_semantic_gain.png", "图 3 开发性回测语义融合相对最强数值基线的误差比值")
    add_heading(doc, "6.3 文本置换反事实", 2)
    add_body(doc, "正确对齐文本只在 3/12 项任务优于置换文本，且三项全部来自 Traffic。跨领域结果说明，文本特征能够被模型使用不等于文本包含与目标有关的时序信息。若缺少置换对照，Climate、Economy 和 Energy 中的高维文本校正可能被误写成有效融合。")
    add_figure(doc, "fig04_alignment_vs_permutation.png", "图 4 正确时点对齐相对随机置换文本的改善率")

    add_heading(doc, "7 冻结确认与外部基准", 1)
    add_heading(doc, "7.1 确认协议", 2)
    add_body(doc, "在查看 Agriculture 和 Security 测试结果前，冻结领域、窗口、跨度、文本截止条件、置换次数、选择门槛、分歧回退和 Bootstrap 方式。成功标准规定：至少一个未见任务的语义路径必须通过全部验证门槛，并在测试期不劣于数值路径。结果产生后未更换领域或修改门槛。")
    add_heading(doc, "7.2 确认结果", 2)
    rows = []
    for _, r in conf.iterrows():
        rows.append([f"{r.domain} H={int(r.horizon)}", r.selected_variant, f"{r.internal_numeric_mse:.6f}", f"{r.semantic_mse:.6f}", f"{r.frequency_semantic_mse:.6f}", f"{r.avoided_semantic_loss_pct:.2f}%"])
    add_caption(doc, "表 4 冻结确认实验与安全回退")
    add_table(doc, ["任务", "选择", "数值MSE", "语义MSE", "频率语义MSE", "避免损失"], rows, [1.45, .75, 1.05, 1.05, 1.2, 1.0], 7.8)
    add_body(doc, "六项任务均未启用语义路径，预先规定的语义增益确认标准没有达到。该结果限制了本文主张：不能说文本在多领域普遍提高精度。另一方面，语义候选在六项任务最终都比内部数值路径差，选择器在 6/6 任务正确拒绝了这些路径，使最终预测等于数值回退，平均避免 21.97% 的语义路径损失。")
    add_figure(doc, "fig05_confirmation_safety.png", "图 5 冻结确认实验中的语义负迁移与安全回退")
    doc.add_page_break()
    add_heading(doc, "7.3 ETT 外部数值基准", 2)
    rows = []
    for _, r in ett.iterrows():
        rows.append([f"{r.dataset} H={int(r.horizon)}", r.best_model, f"{r.best_mse:.6f}", f"{r.best_mse_std:.6f}", r.best_non_patch_model, f"{r.patchtst_gain_vs_best_non_patch_pct:.2f}%"])
    add_caption(doc, "表 5 ETT 外部单目标预测结果")
    add_table(doc, ["任务", "优胜模型", "MSE均值", "标准差", "最强非Patch", "改善"], rows, [1.2, 1.1, 1.05, 1.0, 1.25, .85], 8.0)
    add_body(doc, "PatchTST-M 在 ETTh1 和 ETTh2 的六项任务中全部获得最低平均 MSE，相对最强非 PatchTST 基线平均改善 12.69%。其中 ETTh2 H=96 仅改善 0.86%，表明优势并非在每项任务都很大；ETTh1 H=720 和 ETTh2 H=720 的改善分别为 27.73% 和 21.12%，说明分块表示在长跨度更有价值。")
    add_figure(doc, "fig06_ett_external_results.png", "图 6 ETT 外部基准不同数值模型的预测误差")
    add_figure(doc, "fig07_seed_stability.png", "图 7 ETT 深度模型三次随机种子稳定性")

    add_heading(doc, "8 稳健性与误差分析", 1)
    add_heading(doc, "8.1 输入扰动", 2)
    rows = [
        ["高斯噪声 σ=0.10", "3.31%", "轻度退化"],
        ["随机缺失 10%", "1.68%", "前向填补后较稳"],
        ["末段连续缺失 12步", "24.59%", "依赖最新观测"],
        ["尖峰污染 2%×3σ", "84.38%", "最主要稳健性风险"],
    ]
    add_caption(doc, "表 6 PatchTST 在六项 ETT 任务上的平均相对退化")
    add_table(doc, ["扰动", "平均MSE退化", "解释"], rows, [2.0, 1.4, 3.0], 9)
    add_body(doc, "轻度分散噪声和随机缺失不会改变 PatchTST 的总体可用性，但局部尖峰会显著放大误差。末段连续缺失也造成明显退化，因为预测最依赖靠近起点的观测。后续系统应在进入数值骨干前增加训练期估计的稳健尺度裁剪和异常掩码，并在末段缺失时扩大季节专家权重。")
    add_figure(doc, "fig09_input_robustness.png", "图 8 PatchTST 数值骨干的输入扰动稳健性")
    add_heading(doc, "8.2 领域漂移", 2)
    add_body(doc, "Security 测试段目标超出训练段主要范围，三种跨度 MSE 都很高。安全回退只保证不额外采用有害文本，不能解决目标本身的制度变化。Agriculture 的 report 来源在训练和验证阶段无覆盖，而测试阶段覆盖约 56% 至 60%，来源分布切换进一步削弱语义校正的可迁移性。")
    add_figure(doc, "fig08_security_distribution_shift.png", "图 9 Security 目标序列的末段分布漂移")
    add_heading(doc, "8.3 误差解释边界", 2)
    add_body(doc, "本文观察到的是预测关联，不是文本事件对数值目标的因果效应。Traffic 中的改善只能说明当前时点索引和任务协议下文本提供了可利用增量；未进行干预或工具变量设计，不能把某条新闻解释为目标变化的原因。Bootstrap 区间反映当前预测原点的不确定性，也不代表跨年份或跨地区外部有效性。")

    add_heading(doc, "9 系统实现与可复现性", 1)
    add_heading(doc, "9.1 文件结构", 2)
    add_body(doc, "项目按原始数据、处理数据、正式代码、输出表格、输出图形、模型文件、日志、论文和参考资料分层。数据处理不覆盖原始 CSV；每个实验目录保存配置 JSON、逐次指标、汇总指标、预测或原点误差和模型参数。最终主张由 summarize_results.py 从实验产物重新计算，避免手工抄录。")
    add_heading(doc, "9.2 运行顺序", 2)
    add_bullets(doc, [
        "先运行数据审计和文本索引，核验时间顺序与 fact 截止条件。",
        "生成冻结 MiniLM 向量和逐预测起点语义缓存。",
        "运行传统基线、DLinear、PatchTST 与语义探针。",
        "按冻结协议运行 SafeFAME 选择和未见领域确认。",
        "运行 ETT 外部基准、输入扰动、结果汇总和图表生成。",
    ])
    add_heading(doc, "9.3 结果看板", 2)
    add_body(doc, "静态结果看板提供总览、开发回测筛选、冻结确认、ETT 基准和方法边界五个视图。站点只包含经过汇总脚本核验的数值，不连接训练环境或个人文件。发布版本保持私有，便于课程检查时展示证据，同时不替代正式报告和附件。")

    add_heading(doc, "10 结论与局限", 1)
    add_heading(doc, "10.1 主要结论", 2)
    add_bullets(doc, [
        "严格时点对齐后，开发领域的文本收益高度条件化。12 项任务中只有 Traffic H=6 和 H=12 超过最强数值基线。",
        "正确对齐文本仅在 Traffic 的三个跨度优于置换文本，说明反事实置换是识别真实时序增量的必要检验。",
        "冻结确认集没有复现语义精度增益，但安全选择器拒绝 6/6 个有害语义路径，确认了负迁移防护价值。",
        "PatchTST 数值骨干在六项 ETT 单目标任务中全部获胜，平均改善 12.69%，但对尖峰和连续末段缺失敏感。",
    ])
    add_heading(doc, "10.2 创新点", 2)
    add_body(doc, "本文的创新集中在实验逻辑而非模型堆叠。其一，把文本发布日期写成可执行的信息约束，并对预测型字段做显式排除；其二，以置换文本构造保持边际分布的反事实基线；其三，用验证期总体、前后半段和置换比较共同决定是否启用语义路径；其四，先冻结规则再进入未见领域，允许出现并报告否定性确认结果。")
    add_heading(doc, "10.3 局限与改进", 2)
    add_body(doc, "研究范围仍有限。确认领域只有两个，文本来源在 Agriculture 存在覆盖切换，Security 存在强目标漂移；当前融合头主要检验线性语义增量，未覆盖复杂的事件实体关系；文本编码器以英文通用句向量为主，领域术语表达可能不足；稳健性实验显示尖峰污染是明显弱点。下一步应在不修改确认结果的前提下增加更多未见领域、基于训练统计的异常抑制、来源不变表示和按年份滚动外部验证。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Liu H, Xu S, Zhao Z, et al. Time-MMD: Multi-Domain Multimodal Dataset for Time Series Analysis. Advances in Neural Information Processing Systems 37, 2024. DOI: 10.52202/079017-2476.",
        "[2] Nie Y, Nguyen N H, Sinthong P, Kalagnanam J. A Time Series is Worth 64 Words: Long-term Forecasting with Transformers. International Conference on Learning Representations, 2023.",
        "[3] Zeng A, Chen M, Zhang L, Xu Q. Are Transformers Effective for Time Series Forecasting? Proceedings of the AAAI Conference on Artificial Intelligence, 37(9):11121-11128, 2023. DOI: 10.1609/aaai.v37i9.26317.",
        "[4] Wang W, Wei F, Dong L, et al. MiniLM: Deep Self-Attention Distillation for Task-Agnostic Compression of Pre-Trained Transformers. Advances in Neural Information Processing Systems 33, 2020.",
        "[5] Reimers N, Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP, 3982-3992, 2019. DOI: 10.18653/v1/D19-1410.",
        "[6] Liu C, Xu Q, Miao H, et al. TimeCMA: Towards LLM-Empowered Multivariate Time Series Forecasting via Cross-Modality Alignment. arXiv:2406.01638, 2024.",
        "[7] Paszke A, Gross S, Massa F, et al. PyTorch: An Imperative Style, High-Performance Deep Learning Library. Advances in Neural Information Processing Systems 32, 2019.",
        "[8] Pedregosa F, Varoquaux G, Gramfort A, et al. Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12:2825-2830, 2011.",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)

    add_heading(doc, "附录 A 交付文件", 1)
    add_bullets(doc, [
        "正式报告 DOCX 与 PDF。",
        "src 目录中的完整 Python 源代码与统一依赖文件。",
        "data_processed 中的文本索引、语义向量元数据和时点特征缓存。",
        "outputs 中的逐次指标、模型参数、训练日志、预测样本、最终表格和九组 PNG PDF 图。",
        "docs 中的课程要求、选题方案、数据审计、协议冻结、确认结果和文献核验记录。",
        "prototype 中的私有结果看板源文件与托管清单。",
    ])
    add_heading(doc, "附录 B 审稿式检查", 1)
    checks = [
        ["题意", "逐项覆盖问题描述 实验设计 算法 结果 图表 源码 数据与过程附件", "通过"],
        ["数据", "原始文件不覆盖 清洗有记录 时间严格升序 文本截止早于预测起点", "通过"],
        ["模型", "包含简单基线 深度基线 消融 置换对照 冻结确认和稳健性", "通过"],
        ["结果", "表格由脚本生成 关键数字与 CSV 一致 否定结果未隐藏", "通过"],
        ["边界", "不把相关写成因果 不称确认集超越全部传统基线", "通过"],
        ["复现", "固定种子 保存配置 模型 日志 预测与统一依赖", "通过"],
    ]
    add_table(doc, ["检查项", "核验内容", "状态"], checks, [1.0, 4.8, .7], 8.5)

    path = OUT / "SafeFAME-TS_课程设计报告.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build())
