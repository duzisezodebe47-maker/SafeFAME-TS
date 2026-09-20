"""Add verified v4 results to the retained three-report submission."""
import csv
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from revise_boundary_reports import insert,table_before,font

ROOT=Path(__file__).resolve().parents[1]
BACKUP=ROOT/'paper/archive/20260919_before_v4'
OUT=ROOT/'outputs/safefame_v4'

def picture(anchor,doc,name,caption,width=6.1):
    p=insert(anchor,doc,'');p.paragraph_format.first_line_indent=Pt(0);p.alignment=1
    shape=p.add_run().add_picture(str(OUT/'figures'/name),width=Inches(width))
    shape._inline.docPr.set('descr',caption)
    p.paragraph_format.keep_with_next=True
    insert(anchor,doc,caption,'Caption')

def main():
    summary=json.loads((OUT/'summary.json').read_text(encoding='utf-8'))
    verified=json.loads((OUT/'verification.json').read_text(encoding='utf-8'))
    assert verified['status']=='PASS' and verified['task_folds']==120 and verified['replay_models']
    cfg=json.loads((ROOT/'submission.json').read_text(encoding='utf-8'))
    domains=list(csv.DictReader((OUT/'domain_results.csv').open(encoding='utf-8')))
    data=list(csv.DictReader((ROOT/'data_processed/v4/numerical_audit.csv').open(encoding='utf-8')))
    sem=summary['semantic_residual'];freq=summary['frequency_residual']
    overview=(f'扩展v4完成9个领域、10组序列的40项任务与120次滚动评估，使用5个随机种子。'
              f'完整门控实际启用文本{summary["text_selected"]}次，数值回退{summary["numeric_fallback"]}次；'
              f'启用后测试点改善{summary["selected_point_positive"]}次、变差{summary["selected_point_negative"]}次。'
              '历史v2/v3的18项全部回退结论保留，但不能外推到这一扩大后的协议。')
    protocol=('v4将每组序列设为4个跨度。三折训练/校准/决策/测试末端比例分别为40/50/70/80%、'
              '50/60/80/90%、60/70/90/100%，逐段剔除目标跨界窗口；三折测试目标区间不重叠，'
              '后折使用当时已观测历史。数值专家、残差候选、正则范围和早停设置沿用原框架。'
              '主门控改为999次逐行联合错位，另做999次循环移位敏感性，二者同为999但用途不同；'
              '历史v2/v3主门控仍为99次。5000次移动块Bootstrap始终只用于测试诊断。')
    boundary=('此次扩展是在已有历史数据上预先固定运行配置后开展的滚动回测，不是新增未来时间段的独立确认。'
              'Time-MMD上游快照未更新；Health_US与Health_AFR是同一健康领域的两组序列。'
              '不同跨度和滚动折共享历史，120次评估不等于120个独立数据集；0.025只对应单个任务—折内两候选的校正，'
              '不声称控制全部240条候选的错误率。未知目标窗口排除后，区间块按剩余起点顺序构造，跨缺口不等于连续日历块。发布时间代理、低功效和非因果解释边界继续适用。')
    for report in cfg['reports']:
        doc=Document(BACKUP/report['docx'])
        if report['role']=='final':
            # Update the abstract without replacing the historical evidence.
            keyword=next(p for p in doc.paragraphs if p.text.startswith('关键词'))
            insert(keyword,doc,overview)
            anchor=next(p for p in doc.paragraphs if p.text=='10 模型评价与结论')
            insert(anchor,doc,'9.5 多领域扩展滚动回测','Heading 2')
            insert(anchor,doc,overview+' 本节单列扩大规模后的证据，不以新结果改写旧实验。')
            insert(anchor,doc,'扩展数据来自同一固定Time-MMD快照[1]。共25613条去重fact，其中14281条复用已核验句向量，新增11332条使用相同固定MiniLM版本编码。Environment保留日频；AFR前三条和SocialGood末尾八条无目标观测记录裁去，AFR另有三个无穷目标转为未知。Health_US五个重复日期对应的目标冲突，标为未知并保留原冲突行，排除覆盖未知目标的训练和评价窗口。输入协变量仅历史前向填充加训练中位数，不把填充值当作目标真值。见表13。')
            insert(anchor,doc,'表 13 扩展数据及四档预测跨度','Caption')
            settings=json.loads((ROOT/'configs/safefame_v4.json').read_text(encoding='utf-8'))['domains']
            table_before(anchor,doc,['序列','原始行','清洗后行','时间步','预测跨度'],[
                [r['domain'],r['raw_rows'],r['clean_rows'],{'daily':'日','weekly':'周','monthly':'月'}[settings[r['domain']]['frequency']],
                 '/'.join(str(h) for h in settings[r['domain']]['horizons'])] for r in data])
            insert(anchor,doc,'目标用该折已观测训练值的均值与总体标准差标准化；DLinear的各输入列另用因果填补后的训练统计量标准化，有未知目标时两者未必相同。全部模型输出和评价真值均处于同一目标标准化坐标，不在不同输入尺度之间直接比较损失。')
            insert(anchor,doc,protocol+' 每个任务—折先保存资格选择记录，再计算测试结果；残差和被选中深度回退在测试之前的完整可用窗口重拟合。未选中深度专家的校准权重用于选择证据，不将其称为测试期模型排名。图13展示时间划分。')
            picture(anchor,doc,'rolling_design.png','图 13 v4三折目标隔离的时间划分')
            insert(anchor,doc,'式（5）的加一经验p定义在v4仍保留，重复数改为999，对应分母1000；逐行错位仍在每次校准中重新选择正则，循环移位固定原对齐正则。MSE沿用式（7），移动块损失区间按第6.4节构造，并与直接逐块采样对照。H=1新增任务显式保持预测二维形状，避免一维输出加Last锚点产生错误广播。早期未通过核验的开发批次已归档，未计入本节。')
            insert(anchor,doc,'表 14 v4候选检验与测试诊断汇总','Caption')
            headers=['指标（每类120条）','语义候选','频率候选']
            fields=[('测试点改善','point_positive'),('测试区间为正','ci_positive'),('测试区间为负','ci_negative'),
                    ('主置换p≤0.025','row_pass'),('完整资格通过','eligible'),('循环移位p≤0.025','circular_pass'),
                    ('相对数值残差对照区间正','matched_ci_positive'),('相对数值残差对照区间负','matched_ci_negative')]
            table_before(anchor,doc,headers,[[label,str(sem[k]),str(freq[k])] for label,k in fields])
            insert(anchor,doc,'语义9条与频率2条完整资格通过，合计11条；其中同一任务—折可同时通过两条，最终仅按决策MSE选一条，因此实际启用为10次，不能把路径通过数相加当成任务数。')
            insert(anchor,doc,f'表14以任务—折为统计单位，每类120条。点改善和区间正负均相对该折校准选择的数值回退；最后两行使用同家族纯数值残差对照。语义候选三折平均相对改善为正的任务族为{sem["task_family_mean_positive"]}/40，频率为{freq["task_family_mean_positive"]}/40；这只是40项任务的描述性汇总，不能替代独立性检验。更多候选达到单项p阈值，不意味着必然通过完整门控。')
            picture(anchor,doc,'domain_results.png','图 14 v4分序列候选收益与实际启用结果')
            insert(anchor,doc,f'图14左侧按每组12个任务—折汇总候选相对MSE改善中位数，右侧只统计实际启用。系统在全部120次评估上的平均相对改善为{summary["selected_mean_gain_pct"]:.2f}%，中位数为{summary["selected_median_gain_pct"]:.2f}%；这是逐任务相对比值的等权描述，不是跨序列MSE直接相加。实际启用后的测试损失差区间有{summary["selected_ci_positive"]}次完全为正、{summary["selected_ci_negative"]}次完全为负，说明资格规则并不提供测试期无害保证。各序列见表15，全部逐项结果保存在candidate_results.csv和task_results.csv。')
            insert(anchor,doc,'表 15 v4分序列实际路由与样本量','Caption')
            table_before(anchor,doc,['序列（各12项）','启用','点改善/变差','决策块中位数','测试文本覆盖'],[
                [r['domain'],r['selected'],r['selected_positive']+'/'+r['selected_negative'],f'{float(r["median_decision_blocks"]):.1f}',f'{float(r["mean_text_coverage_pct"]):.1f}%'] for r in domains])
            insert(anchor,doc,'两次实际点变差分别为Climate H4第3折（−3.01%）和Environment H1第2折（−0.08%），两者区间均跨零；不能称已证明有害，也不能因不显著而写成无风险。Agriculture H12第1折频率候选改善52.95%，是较大的单项贡献；全部120次等权平均0.77%并不代表每个领域或未来时段均有类似收益。')
            insert(anchor,doc,f'共{summary["calibration_windows"]}个校准、{summary["decision_windows"]}个决策和{summary["test_windows"]}个测试起点，这些计数跨跨度重复观测，不能视作独立样本总量。决策块中位数为{summary["median_decision_blocks"]:.1f}，{summary["decision_blocks_ge8"]}/120次达到8块，{summary["decision_blocks_lt3"]}/120次不足3块。图15的8块线只是参考；决策窗口与前后折依然依赖，低频长跨度的有效信息仍有限，不能用增加重采样次数替代新增观测。')
            insert(anchor,doc,f'另有{len(summary["zero_train_text_tasks"])}个任务—折的训练起点文本覆盖为零，详见task_results.csv。零向量导致PCA解释方差比未定义，但投影、预测和保存指标均通过有限性检查；不使用该方差比作任何选择。这类任务缺乏可学习的训练文本变化，即使后期文本出现，也不能将结果解释为一般文本价值检验。')
            picture(anchor,doc,'decision_blocks.png','图 15 v4决策期完整非重叠块诊断')
            insert(anchor,doc,boundary+' v3与v4同时改变领域、跨度、划分宽度、种子和重复次数，差异不能单独归因为某一项。端到端已保存1200个深度校准检查点、逐种子日志、逐起点预测、两种零分布和选中专家重拟合权重，核验从原始观测与预测重算路由、指标和区间，并回放选中深度权重；不代表重新执行全部置换拟合。')
            conclusion=next(p for p in doc.paragraphs if p.text=='10.2 优点')
            insert(conclusion,doc,overview+' 扩展结果表明结论依赖任务与协议，最终判断应同时查看实际收益、变差案例、数值残差对照和不确定性。')
            appendix=next(p for p in doc.paragraphs if p.text.replace(' ','')=='附录C资格门控算法伪代码')
            appendix.paragraph_format.page_break_before=True
            note=doc.add_paragraph('版本口径：以下99次伪代码保留历史v2/v3设置。v4对每个滚动折重复同类选择流程，主逐行错位改为999次；折边界按第9.5节，目标跨界窗口剔除。附加999次循环移位和5000次测试区间仍只作诊断。')
            font(note);note.paragraph_format.first_line_indent=Pt(22)
            appendix._p.addnext(note._p)
            for p in doc.paragraphs:
                if p.text.startswith('下一步应优先获得可核验发布时间'):
                    p.text=('v4已完成扩大领域、三折滚动与五种子训练，但共享历史仍不构成独立确认。后续应优先核验真实发布时间，预先固定7天、30天、60天发布滞后敏感性；在真正新增时间段上锁定同一规则作确认。针对未知目标造成的缺口，按连续日历块复核区间；针对Last锚点不对称，另设围绕已选数值专家的文本残差对照。预先定义误启用风险与收益的权衡，不根据当前测试结果放宽门槛。')
                    font(p)
                elif p.text.startswith('模型直接回答何时启用文本'):
                    p.text=p.text.replace('三随机种子','历史三种子及v4五种子');font(p)
                elif p.text.startswith('频率语义交互没有形成'):
                    p.text='历史v2/v3中，'+p.text;font(p)
            for t in doc.tables:
                for row in t.rows:
                    if len(row.cells)==3 and row.cells[0].text in ('系统或实验设计','实验结果及分析图表'):
                        row.cells[1].text='第5至6章及第9.5节' if row.cells[0].text=='系统或实验设计' else '第7至9章'
                        for p in row.cells[1].paragraphs:
                            font(p)
                            for run in p.runs:run.font.size=Pt(8)
            for t in doc.tables:
                if len(t.rows)==1 and len(t.columns)==1 and '算法1' in t.cell(0,0).text.replace(' ',''):
                    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        else:
            anchor=next(p for p in doc.paragraphs if p.text=='参考文献')
            insert(anchor,doc,'扩展实验完成后的补记','Heading 2')
            insert(anchor,doc,'本节在扩大实验完成后补充，用于与最终报告保持一致，不表示第4周或第10周当时已完成全部扩展实验。'+overview)
            insert(anchor,doc,protocol)
            insert(anchor,doc,f'扩展语义候选在{sem["point_positive"]}/120次测试中点改善，区间{sem["ci_positive"]}正{sem["ci_negative"]}负；频率候选在{freq["point_positive"]}/120次点改善。决策块中位数为{summary["median_decision_blocks"]:.1f}，{summary["decision_blocks_ge8"]}/120次达到8块参考值。原始记录未覆盖，全部结果、源代码和核验记录在outputs/safefame_v4，正式方法与结果见最终报告第9.5节。')
            insert(anchor,doc,boundary)
        doc.core_properties.comments='v4 expanded rolling backtest; historical v2/v3 retained; not prospective confirmation'
        doc.save(ROOT/report['docx'])
    print('Three reports updated from verified v4 outputs; render review required')

if __name__=='__main__':main()
