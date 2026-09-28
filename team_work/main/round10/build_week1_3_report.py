from __future__ import annotations

import json
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\aaa2026人工智能课程报告")
REPO = ROOT / "SafeFAME-TS_主控技术工作区"
OUT = ROOT / "1-3周作业提交"
ASSETS = OUT / "图表素材"
SCORE = REPO / "team_work" / "main" / "round10" / "results" / "Climate_test_score.json"
DOCX = OUT / "SafeFAME-TS_第1-3周选题调研与初步技术方案_提交版.docx"

NAVY = "17365D"
BLUE = "2F75B5"
SKY = "DDEBF7"
LIGHT = "F3F6FA"
ORANGE = "ED7D31"
GREEN = "70AD47"
RED = "C00000"
GRAY = "666666"
WHITE = "FFFFFF"

FONT_CN = r"C:\Windows\Fonts\msyh.ttc"
FONT_CN_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"


def font(size: int, bold: bool = False):
    path = FONT_CN_BOLD if bold and Path(FONT_CN_BOLD).exists() else FONT_CN
    return ImageFont.truetype(path, size)


def rounded(draw, xy, radius, fill, outline=None, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_text(draw, xy, text, fnt, fill):
    x1, y1, x2, y2 = xy
    box = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=8, align="center")
    w, h = box[2] - box[0], box[3] - box[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=fnt, fill=fill, spacing=8, align="center")


def make_route_figure(path: Path):
    img = Image.new("RGB", (1800, 980), "#F5F8FC")
    d = ImageDraw.Draw(img)
    d.text((90, 48), "SafeFAME-TS 改进技术路线", font=font(58, True), fill="#17365D")
    d.text((92, 122), "以时间切分和安全门控约束多模态增量，测试集只在路线冻结后打开一次", font=font(28), fill="#55657A")
    boxes = [
        (70, 240, 350, 450, "原始时间序列\n文本与质量字段", "#DDEBF7"),
        (420, 240, 700, 450, "时间审计与切分\n训练 / 选择 / 测试", "#E2F0D9"),
        (770, 240, 1050, 450, "候选模型注册\n基线 / N / N+S+Q", "#FFF2CC"),
        (1120, 240, 1400, 450, "选择区安全门控\n性能 + 稳定性 + 置换", "#FCE4D6"),
        (1470, 240, 1730, 450, "冻结路线\n一次性测试", "#E4DFEC"),
    ]
    for x1, y1, x2, y2, label, color in boxes:
        rounded(d, (x1, y1, x2, y2), 24, color, "#A7B4C5", 3)
        centered_text(d, (x1 + 12, y1 + 10, x2 - 12, y2 - 10), label, font(33, True), "#17365D")
    for i in range(len(boxes) - 1):
        x = boxes[i][2] + 15
        y = 345
        d.line((x, y, x + 40, y), fill="#2F75B5", width=8)
        d.polygon([(x + 40, y), (x + 22, y - 12), (x + 22, y + 12)], fill="#2F75B5")
    stages = [
        (110, 585, 510, 830, "防泄漏", "所有归一化、特征学习和调参\n均限制在当时可见的数据内"),
        (700, 585, 1100, 830, "防负迁移", "多模态候选只有在选择区同时满足\n性能与稳定性门槛才可进入测试"),
        (1290, 585, 1690, 830, "可证伪", "保留朴素基线；若语义信息无效，\n系统自动退化为数值路线"),
    ]
    for x1, y1, x2, y2, title, body in stages:
        rounded(d, (x1, y1, x2, y2), 22, "#FFFFFF", "#2F75B5", 3)
        d.text((x1 + 30, y1 + 32), title, font=font(34, True), fill="#2F75B5")
        d.multiline_text((x1 + 30, y1 + 98), body, font=font(25), fill="#333333", spacing=12)
    img.save(path, quality=95)


def make_domain_figure(path: Path):
    img = Image.new("RGB", (1800, 920), "#F7F9FC")
    d = ImageDraw.Draw(img)
    d.text((85, 48), "四领域渐进验证设计", font=font(58, True), fill="#17365D")
    d.text((88, 122), "先在中小规模任务验证协议，再扩展到文本稀疏与大规模场景", font=font(28), fill="#55657A")
    items = [
        ("01", "Agriculture", "h=12, season=12", "已完成闭环", "#70AD47"),
        ("02", "Climate", "h=4, season=12", "已完成闭环", "#70AD47"),
        ("03", "SocialGood", "h=3, season=12", "待数据验收", "#ED7D31"),
        ("04", "Environment", "h=7, season=7", "待启动", "#7F8C8D"),
    ]
    x_positions = [85, 515, 945, 1375]
    for x, (idx, name, spec, status, color) in zip(x_positions, items):
        rounded(d, (x, 255, x + 340, 700), 28, "#FFFFFF", color, 5)
        d.ellipse((x + 115, 290, x + 225, 400), fill=color)
        centered_text(d, (x + 115, 290, x + 225, 400), idx, font(42, True), "#FFFFFF")
        centered_text(d, (x + 20, 430, x + 320, 500), name, font(34, True), "#17365D")
        centered_text(d, (x + 20, 510, x + 320, 565), spec, font(25), "#4D5A6A")
        rounded(d, (x + 70, 605, x + 270, 665), 18, color, color, 1)
        centered_text(d, (x + 70, 605, x + 270, 665), status, font(24, True), "#FFFFFF")
    d.line((255, 760, 1545, 760), fill="#AAB7C8", width=8)
    d.polygon([(1585, 760), (1535, 735), (1535, 785)], fill="#AAB7C8")
    d.text((650, 805), "任务难度与工程规模逐步增加", font=font(28, True), fill="#55657A")
    img.save(path, quality=95)


def make_climate_chart(path: Path, score: dict):
    metrics = score["metrics_standardized_and_raw"]
    labels = ["N", "Last", "AR-Ridge", "SeasonalNaive"]
    vals = [metrics[x]["mse"] for x in labels]
    img = Image.new("RGB", (1700, 1000), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((90, 55), "Climate 一次性测试：标准化 MSE", font=font(54, True), fill="#17365D")
    d.text((90, 125), "124 个预测起点 × 4 步；路线在查看测试得分前已冻结", font=font(26), fill="#55657A")
    x0, y0, x1, y1 = 170, 850, 1600, 220
    d.line((x0, y0, x1, y0), fill="#59697C", width=4)
    maxv = max(vals) * 1.08
    colors = ["#2F75B5", "#A5A5A5", "#ED7D31", "#70AD47"]
    bw, gap = 220, 100
    for i, (lab, val, color) in enumerate(zip(labels, vals, colors)):
        x = x0 + 90 + i * (bw + gap)
        h = (val / maxv) * (y0 - y1)
        d.rectangle((x, y0 - h, x + bw, y0), fill=color)
        d.text((x + 20, y0 - h - 55), f"{val:.3f}", font=font(28, True), fill="#17365D")
        centered_text(d, (x - 20, y0 + 20, x + bw + 20, y0 + 85), lab, font(25, True), "#333333")
    d.text((92, 925), "注：SeasonalNaive 数值较大，使前三个柱形视觉差距被压缩；具体比较以表格和置信区间为准。", font=font(22), fill="#666666")
    img.save(path, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            tag = "w:" + edge
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn("w:" + key), str(value))


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(p, first_line=True, line=20, before=0, after=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing = Pt(line)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if first_line:
        pf.first_line_indent = Pt(21)


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, "宋体", 10.5, True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph(p, first_line=False, line=19, after=2)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = level == 1
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=False, line=18, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_run_font(r, "宋体", 9.5, False, GRAY)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths=None, note=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], NAVY)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, "微软雅黑", 9, True, WHITE)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cells[i], LIGHT)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(16)
                for r in p.runs:
                    set_run_font(r, "宋体", 8.5)
    border = {"val": "single", "sz": "4", "color": "B7C9DD"}
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, cell in enumerate(row.cells):
            if widths:
                cell.width = Cm(widths[i])
            set_cell_border(cell, top=border, bottom=border, left=border, right=border)
    if note:
        p = doc.add_paragraph()
        set_paragraph(p, first_line=False, line=16, after=5)
        r = p.add_run(note)
        set_run_font(r, "宋体", 8.5, False, GRAY)
    return table


def add_formula(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, first_line=False, line=20, before=4, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(text)
    set_run_font(r, "Cambria Math", 11, False, NAVY)


def add_picture(doc, path, width_cm=16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.extend([fldChar1, instrText, fldChar2])


def build_doc():
    OUT.mkdir(parents=True, exist_ok=True)
    ASSETS.mkdir(parents=True, exist_ok=True)
    score = json.loads(SCORE.read_text(encoding="utf-8"))
    route_fig = ASSETS / "图1_改进技术路线.png"
    domain_fig = ASSETS / "图2_四领域渐进验证.png"
    climate_fig = ASSETS / "图3_Climate测试MSE.png"
    make_route_figure(route_fig)
    make_domain_figure(domain_fig)
    make_climate_chart(climate_fig, score)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.gutter = Cm(0.5)
    sec.header_distance = Cm(1.7)
    sec.footer_distance = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for level, size in [(1, 18), (2, 15), (3, 12)]:
        st = styles[f"Heading {level}"]
        st.font.name = "黑体"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
        st.paragraph_format.space_before = Pt(12 if level == 1 else 8)
        st.paragraph_format.space_after = Pt(8 if level == 1 else 4)
        st.paragraph_format.line_spacing = Pt(22)

    for section in doc.sections:
        h = section.header.paragraphs[0]
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = h.add_run("人工智能课程设计｜SafeFAME-TS")
        set_run_font(r, "微软雅黑", 8.5, False, GRAY)
        add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("人工智能课程设计")
    set_run_font(r, "黑体", 22, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("SafeFAME-TS")
    set_run_font(r, "Arial", 30, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向多领域多模态时间序列预测的安全门控框架")
    set_run_font(r, "黑体", 20, True, BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("第 1—3 周选题调研与初步技术方案")
    set_run_font(r, "微软雅黑", 15, False, ORANGE)
    doc.add_paragraph("\n")
    cover_rows = [
        ("课程名称", "人工智能课程设计"),
        ("组别", "三人小组"),
        ("成员 1", "姓名：____________    学号：____________"),
        ("成员 2", "姓名：____________    学号：____________"),
        ("成员 3", "姓名：____________    学号：____________"),
        ("提交日期", "2026 年 9 月 28 日"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for left, right in cover_rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = left, right
        cells[0].width, cells[1].width = Cm(4.0), Cm(10.0)
        for i, c in enumerate(cells):
            set_cell_shading(c, SKY if i == 0 else WHITE)
            for pp in c.paragraphs:
                pp.paragraph_format.space_before = pp.paragraph_format.space_after = Pt(5)
                for rr in pp.runs:
                    set_run_font(rr, "宋体", 11, i == 0, NAVY if i == 0 else None)
            set_cell_border(c, bottom={"val": "single", "sz": "6", "color": "B7C9DD"})
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("说明：成员姓名与学号由小组提交前填写。")
    set_run_font(r, "宋体", 9, False, GRAY)

    # Abstract
    add_heading(doc, "摘要", 1)
    add_body(doc, "本项目研究多领域时间序列预测中“外部文本和质量信息何时真正有用”的问题。公开多模态数据集为农业、气候、社会公共事务和环境等领域同时提供数值序列、文本描述与数据质量字段，但简单拼接多种模态可能造成时间泄漏、噪声放大和负迁移。为此，小组拟构建 SafeFAME-TS：一个以时间因果切分、候选模型注册、安全门控和一次性测试为核心的可复核框架。框架先建立 Last、SeasonalNaive 和 AR-Ridge 等朴素基线，再比较纯数值候选 N、质量感知候选 N+Q、语义融合候选 N+S+Q 及其稳定性增强版本 N+S+Q+SF。只有当多模态候选在选择区同时满足误差改善、前后半段稳定和时间对齐置换检验时，才允许进入冻结测试；否则自动退化为更简单的数值路线。")
    add_body(doc, "第 1—3 周完成了问题定义、论文与开源项目调研、数据任务拆解、关键算法选择、改进路线设计和实验协议冻结。当前工程预实验已在 Agriculture 与 Climate 两个领域完成从候选选择到一次性测试的闭环。Climate 任务上，冻结候选 N 的测试 MSE 为 0.1340，低于 Last 的 0.1446，但二者损失差的移动区组 Bootstrap 95% 区间跨过 0，因此本项目不会把该结果夸大为显著优势。后续将在 SocialGood 与 Environment 上继续验证安全退化和跨领域稳定性。")
    p = doc.add_paragraph()
    set_paragraph(p, first_line=False, line=20)
    r = p.add_run("关键词：")
    set_run_font(r, "黑体", 10.5, True)
    r = p.add_run("多模态时间序列；安全门控；防信息泄漏；负迁移；时间置换检验；可复现评估")
    set_run_font(r)

    add_heading(doc, "1 选题背景与问题介绍", 1)
    add_heading(doc, "1.1 研究背景", 2)
    add_body(doc, "现实中的时间序列并不孤立存在。农业产量同时受到天气与政策信息影响，气候观测伴随预警文本和数据质量记录，社会公共事务常含事件描述，环境监测也包含告警和维护信息。Time-MMD 等公开工作把这些异构信息按时间组织，为研究“文本能否帮助数值预测”提供了基础。然而，多模态信息的存在并不等于其具有预测价值：文本可能稀疏、滞后或只是对当前数值的重复描述，若在切分、标准化或特征学习时处理不严，还会无意中使用未来信息。")
    add_body(doc, "课程设计的核心问题不是再堆叠一个参数量更大的模型，而是建立一套能够拒绝无效复杂度的实验协议：在不知道测试结果的条件下，系统能否判断语义和质量信息值得使用；如果证据不足，能否安全退化到数值基线；最终结论能否被另一台电脑在固定文件和固定随机种子下重复得到。")
    add_heading(doc, "1.2 拟解决的三个研究问题", 2)
    add_bullet(doc, "RQ1：在严格时间隔离下，文本语义和质量字段是否能稳定降低多步预测误差？")
    add_bullet(doc, "RQ2：如何把“性能提升”与“偶然时间对齐”区分开，并防止负迁移候选进入测试？")
    add_bullet(doc, "RQ3：当文本覆盖率为零或极低时，系统能否给出可解释、可审计的安全退化结果？")
    add_heading(doc, "1.3 研究边界", 2)
    add_body(doc, "本项目研究公开数据中的预测协议和工程可复现性，不声称由相关性得到因果结论，也不把四个公开任务推广为所有领域的普遍规律。测试集只用于路线冻结后的一次性评价；任何候选选择、特征筛选和阈值调整都必须在训练区或选择区完成。")

    add_heading(doc, "2 相关论文、开源项目与方法调研", 1)
    add_body(doc, "调研按照“数据集—时间序列骨干—语言模型融合—统计验证—工程复现”五条线开展。Time-MMD 提供了多领域、多模态数据与任务背景；PatchTST、N-BEATS 和 DLinear 展示了强数值基线的重要性；Time-LLM 代表将语言模型用于时间序列的探索；Sentence-BERT 和 MiniLM 为轻量文本表征提供可复用工具；Politis 与 Romano 的平稳 Bootstrap 为相关时间序列上的区间估计提供了经典依据。")
    rows = [
        ("Time-MMD", "多领域多模态数据集", "统一任务来源和字段审计", "公开数据仍需重新核对时间和覆盖率"),
        ("PatchTST", "分块 Transformer", "强数值序列对照", "高复杂度不必然适合小样本"),
        ("N-BEATS", "可解释深度预测骨干", "多步预测结构参考", "训练成本高于朴素基线"),
        ("DLinear", "线性分解基线", "提醒先验证简单模型", "不能覆盖全部非线性场景"),
        ("Time-LLM", "语言模型重编程", "多模态融合思路", "可能引入算力和泄漏风险"),
        ("Sentence-BERT / MiniLM", "轻量句向量", "文本编码备选", "文本稀疏时应安全退化"),
        ("Stationary Bootstrap", "相关样本重采样", "估计损失差区间", "区组长度选择影响区间"),
    ]
    add_table(doc, ["资料", "核心内容", "本项目用途", "使用边界"], rows, [3.2, 3.6, 4.4, 4.6], "表 2-1 相关工作调研矩阵。资料来源见文末参考文献，访问日期均为 2026-09-28。")
    add_heading(doc, "2.1 调研得到的关键判断", 2)
    add_body(doc, "第一，深度模型不是默认正确答案，朴素 Last、季节性朴素和线性自回归必须作为最低比较标准。第二，多模态模型的提升需要在时间保持的切分上验证，随机打乱样本会低估部署难度。第三，文本信息必须经过覆盖率、时间戳和来源审计；不存在可用文本时，模型名称和实际输入必须同步退化。第四，单次平均误差不足以说明稳定性，还要报告时间分段、预测步长和相关样本重采样区间。")

    add_heading(doc, "3 数据对象与任务设计", 1)
    add_picture(doc, domain_fig)
    add_caption(doc, "图 3-1 四领域渐进验证设计。绿色为已闭环，橙色为待数据验收，灰色为待启动。")
    add_body(doc, "项目采用四个领域任务逐步扩展工程规模。Agriculture 与 Climate 用于验证从候选注册、路线冻结到测试评分的完整链路；SocialGood 专门检验文本覆盖不足时的安全退化；Environment 数据规模最大，用于观察协议在高计算量场景下的可执行性。不同领域不强行合并为一个总体，而是分别评价，再讨论协议层面的共同现象。")
    rows = [
        ("Agriculture_h12_f1", "12", "12", "已闭环", "长预测跨度与季节性"),
        ("Climate_h4_f2", "4", "12", "已闭环", "强 Last 基线与短期预测"),
        ("SocialGood_h3_f1", "3", "12", "待验收", "文本可能为零覆盖"),
        ("Environment_h7_f2", "7", "7", "待启动", "大规模计算与环境时序"),
    ]
    add_table(doc, ["任务", "预测步长 h", "周期", "当前状态", "主要检验点"], rows, [4.0, 2.5, 2.2, 2.7, 5.0], "表 3-1 任务设置。状态反映截至 2026-09-28 的工程进度，不代表全部研究已经完成。")
    add_heading(doc, "3.1 数据审计项目", 2)
    add_bullet(doc, "时间范围、采样间隔、重复时间戳和非单调记录；")
    add_bullet(doc, "数值、文本和质量字段在训练区、选择区、测试区的覆盖率；")
    add_bullet(doc, "缺失值和异常值处理前后的行数变化，原始字段与修正字段并存；")
    add_bullet(doc, "监督窗口是否跨越冻结边界，输入区间和预测区间是否重叠；")
    add_bullet(doc, "每个数据包的逐文件 SHA-256、生成环境和可复算日志。")

    add_heading(doc, "4 关键算法与模型候选", 1)
    add_heading(doc, "4.1 监督窗口与评价指标", 2)
    add_body(doc, "设长度为 L 的历史窗口为 xₜ=[yₜ₋L₊₁,…,yₜ]，模型输出未来 h 步预测 ŷₜ₊₁:ₜ₊ₕ。所有标准化参数只用当时训练区估计。主评价指标为均方误差，同时报告平均绝对误差和均方根误差：")
    add_formula(doc, "MSE = (1 / nh) Σᵢ Σⱼ (yᵢⱼ − ŷᵢⱼ)²    MAE = (1 / nh) Σᵢ Σⱼ |yᵢⱼ − ŷᵢⱼ|")
    add_heading(doc, "4.2 基线与候选注册", 2)
    rows = [
        ("Last", "最后一个观测值向前复制", "最强短期朴素基线之一"),
        ("SeasonalNaive", "复制上一个季节对应值", "检验固定周期是否足够"),
        ("AR-Ridge", "滞后特征 + 岭回归", "低方差数值基线"),
        ("N", "数值残差与冻结正则化网格", "默认数值候选"),
        ("N+Q", "N + 数据质量特征", "检验质量元数据增益"),
        ("N+S+Q", "数值 + 语义 + 质量", "核心多模态候选"),
        ("N+S+Q+SF", "加入稳定性特征/筛选", "降低时段漂移风险"),
    ]
    add_table(doc, ["模型", "输入或结构", "角色"], rows, [3.2, 7.2, 6.0], "表 4-1 候选模型注册表。最终名称必须与实际使用的输入一致。")
    add_body(doc, "岭回归候选通过以下目标函数估计参数，其中 α 只能从预先冻结的网格 {0.01, 0.1, 1, 10, 100, 1000, 10000} 中选择：")
    add_formula(doc, "β̂ = arg minβ [ Σᵢ (yᵢ − β₀ − xᵢᵀβ)² + α‖β‖₂² ]")
    add_heading(doc, "4.3 文本与质量特征", 2)
    add_body(doc, "文本先进行时间戳和覆盖率审计，再使用冻结句向量编码器产生低维语义表征；质量特征包括缺失标记、异常标记和来源可靠性字段。编码器不得在测试文本上继续训练。若训练区文本覆盖为零，语义候选必须退化为已登记的数值或质量感知候选，并在清单中如实记录实际特征，禁止保留具有误导性的“多模态”名称。")

    add_heading(doc, "5 改进技术路线：安全门控与一次性测试", 1)
    add_picture(doc, route_fig)
    add_caption(doc, "图 5-1 SafeFAME-TS 改进技术路线。图中门控仅使用训练区和选择区证据。")
    add_heading(doc, "5.1 选择区门控规则", 2)
    add_body(doc, "候选模型只有同时满足三类条件才可被选中：第一，相对当前回退模型的选择区 MSE 有正向改善；第二，在选择区前半段和后半段均不出现方向反转；第三，在保持数值序列不变、打乱文本与时间对应关系的置换实验中，真实对齐的改善足够罕见。若任一条件失败，系统回退到已经登记的简单候选。")
    add_formula(doc, "p_perm = [1 + Σᵦ I(Δᵦ ≥ Δ_real)] / (B + 1),    B = 999,    门槛 p_perm ≤ 0.025")
    add_body(doc, "这里的行置换检验是决策门控，循环位移只作为诊断，不替代正式门槛。测试区在候选、正则化参数、回退路线和全部哈希冻结后仅打开一次；测试结果不得用于重新选择模型。")
    add_heading(doc, "5.2 该路线相对普通实验的改进", 2)
    add_bullet(doc, "把“融合后平均误差更低”改为可执行的多条件门控，主动拒绝不稳定增益；")
    add_bullet(doc, "把数据、路线和预测文件都绑定到哈希，使报告结论能追溯到具体字节；")
    add_bullet(doc, "将模型电脑与主控评分职责分离，模型电脑提交预测，主控电脑持有测试真值并评分；")
    add_bullet(doc, "把无文本或低覆盖情况写成正式退化契约，避免系统在输入不成立时仍声称使用多模态。")

    add_heading(doc, "6 实验设计与验证方案", 1)
    add_heading(doc, "6.1 时间切分和防泄漏", 2)
    add_body(doc, "每个任务按时间顺序划分训练区、选择区和测试区。监督窗口必须完全落在相应可见范围内；标准化、特征选择、文本编码器拟合和 α 选择都在训练区内部完成。选择区用于候选门控，测试区仅用于最终一次性评分。随机划分不作为正式结果，因为它会破坏时间部署条件并可能让相邻窗口泄漏。")
    add_heading(doc, "6.2 对比、消融与稳健性", 2)
    rows = [
        ("基线比较", "Last、SeasonalNaive、AR-Ridge", "复杂模型是否真正必要"),
        ("模态消融", "N、N+Q、N+S+Q、N+S+Q+SF", "语义、质量和稳定性模块的增量"),
        ("时间分段", "选择区/测试区前半与后半", "提升是否集中在单一时段"),
        ("置换检验", "文本行置换 999 次", "真实时间对齐是否优于偶然对齐"),
        ("区组 Bootstrap", "按有序预测起点重采样", "相关误差下的损失差区间"),
        ("预测步长", "分别报告第 1…h 步 MSE", "误差是否随步长快速放大"),
        ("安全退化", "零文本、稀疏文本、错位文本", "输入失效时是否可控"),
    ]
    add_table(doc, ["实验", "设置", "回答的问题"], rows, [3.0, 6.4, 6.8], "表 6-1 实验矩阵。所有随机过程固定种子，并保存逐次结果或可复算日志。")
    add_heading(doc, "6.3 完成判据", 2)
    add_body(doc, "一个领域只有在数据包验收、候选注册、选择区门控、路线冻结、测试预测、主控评分和结果清单全部通过后才算闭环。若测试点估计改善但区间跨过 0，应报告“点估计更优但稳定证据不足”；若候选在选择区失败，回退本身是合格结果，而不是实验失败。")

    add_heading(doc, "7 当前工程预实验与初步发现", 1)
    add_body(doc, "以下结果用于证明技术路线可执行，并帮助确定后续实验重点。它们来自已冻结路线的一次性测试，不用于反向修改 Climate 候选。")
    add_picture(doc, climate_fig)
    add_caption(doc, "图 7-1 Climate 一次性测试标准化 MSE。数据来源：主控评分结果 Climate_test_score.json。")
    m = score["metrics_standardized_and_raw"]
    rows = [(k, f"{m[k]['mse']:.6f}", f"{m[k]['mae']:.6f}", f"{m[k]['rmse']:.6f}") for k in ["N", "Last", "AR-Ridge", "SeasonalNaive"]]
    add_table(doc, ["方法", "MSE", "MAE", "RMSE"], rows, [4.0, 4.0, 4.0, 4.0], "表 7-1 Climate 测试指标。样本为 124 个预测起点、每个起点 4 步，共 496 行预测。")
    add_body(doc, "N 的 MSE 为 0.134022，比 Last 低 7.33%，比 AR-Ridge 低 17.92%。前半段和后半段相对 Last 的点估计提升分别为 9.20% 和 2.02%，方向一致。但按有序预测起点进行移动区组 Bootstrap 后，N 与 Last 的平均起点 MSE 差 95% 区间为 [−0.00654, 0.02784]，跨过 0。由此只能认为 N 在本次测试上的点估计最低，不能声称其相对 Last 已达到稳定显著优势。")
    add_body(doc, "Agriculture 已完成另一套完整闭环，所选多模态稳定性候选在冻结测试上明显优于 AR-Ridge，为继续研究质量与语义信息提供了积极证据。但跨领域结论仍需 SocialGood 和 Environment 完成后才能形成；当前结果不支持“多模态在所有任务上都更好”的普遍断言。")

    add_heading(doc, "8 创新点、可行性与风险控制", 1)
    add_heading(doc, "8.1 预期创新点", 2)
    add_bullet(doc, "选择机制创新：把多模态增益拆成性能、时间稳定和时间对齐三类证据，并形成可执行门控。")
    add_bullet(doc, "安全退化创新：把零文本或低覆盖处理从临时代码分支提升为可测试契约，模型名称与实际输入同步变化。")
    add_bullet(doc, "证据链创新：以数据包、路线、预测和评分清单的多级哈希保证结果可追溯。")
    add_bullet(doc, "跨领域验证：不合并口径不同的数据，而是比较同一协议在四类任务上的成功、回退与失败模式。")
    add_heading(doc, "8.2 可行性", 2)
    add_body(doc, "项目已有可运行仓库、固定任务规格、三台电脑分工和两个领域闭环结果。基础候选以线性与轻量特征为主，能够在课程周期内完成；复杂语义模型只在覆盖率和门控条件满足时启用。所有核心输出均为 CSV、JSON、PNG 和日志，可由统一入口重新生成。")
    add_heading(doc, "8.3 主要风险与应对", 2)
    rows = [
        ("文本覆盖不足", "语义候选无法训练或标签失真", "执行安全退化并保留覆盖率证据"),
        ("时间泄漏", "测试指标虚高", "严格时间切分，预处理只在训练折拟合"),
        ("负迁移", "多模态反而降低性能", "门控失败即回退，不强行融合"),
        ("小样本波动", "平均误差不稳定", "时间分段、区组 Bootstrap 与基线比较"),
        ("工程不可复现", "队友环境结果不一致", "哈希清单、固定种子、干净目录重放"),
        ("算力超期", "Environment 无法按期完成", "先跑低成本基线，再按门控逐级增加候选"),
    ]
    add_table(doc, ["风险", "可能后果", "控制措施"], rows, [3.4, 5.0, 7.8], "表 8-1 风险矩阵。风险控制是模型设计的一部分，而非报告末尾的形式化说明。")

    add_heading(doc, "9 三人分工与实施计划", 1)
    add_heading(doc, "9.1 小组工程分工", 2)
    rows = [
        ("成员 1（主控）", "协议冻结、独立评分、跨分支验收、集成与报告", "门控 JSON、测试评分、主分支、最终报告"),
        ("成员 2（数据）", "数据审计、窗口构造、切分冻结、哈希打包", "data_processed、split、coverage、manifest"),
        ("成员 3（模型）", "候选实现、选择区实验、预测交付、单元测试", "模型代码、选择证据、预测 CSV、测试日志"),
    ]
    add_table(doc, ["角色", "负责范围", "主要交付"], rows, [3.4, 6.6, 6.2], "表 9-1 三人分工。成员姓名和学号在封面填写，角色边界用于减少测试信息交叉。")
    add_heading(doc, "9.2 周次计划", 2)
    rows = [
        ("第 1 周", "完成选题、问题边界、数据与论文检索", "调研矩阵、研究问题"),
        ("第 2 周", "复现基线并设计防泄漏切分", "基线代码、数据审计规则"),
        ("第 3 周", "冻结技术路线、门控规则和实验矩阵", "本报告、任务规格"),
        ("第 4—6 周", "完成 Agriculture 与 Climate 闭环", "选择证据、测试评分"),
        ("第 7—9 周", "完成 SocialGood，启动 Environment", "退化契约、领域结果"),
        ("第 10 周", "中期汇报与异常修正", "PPT、阶段验收表"),
        ("第 11—15 周", "完成 Environment、消融与稳健性", "全部图表和结果"),
        ("第 16—17 周", "报告定稿、代码复现和答辩准备", "DOCX/PDF、源码、演示"),
    ]
    add_table(doc, ["阶段", "主要工作", "可验收产物"], rows, [3.0, 7.6, 5.6], "表 9-2 实施计划。后续进度以数据包实际质量和门控结果为准，不预设多模态一定获胜。")

    add_heading(doc, "10 课程设计报告主体大纲", 1)
    outline = [
        ("第 1 章 绪论", "背景、问题、研究价值和贡献边界"),
        ("第 2 章 相关工作", "多模态数据集、数值预测骨干、语言模型融合与稳健评估"),
        ("第 3 章 数据与任务", "四领域数据审计、时间切分、字段覆盖和任务定义"),
        ("第 4 章 SafeFAME-TS 方法", "候选注册、特征模块、安全门控、退化契约和哈希证据链"),
        ("第 5 章 实验设计", "基线、消融、置换、Bootstrap、分段与预测步长分析"),
        ("第 6 章 结果与讨论", "各领域结果、跨领域共同点、失败模式与适用边界"),
        ("第 7 章 系统实现", "仓库结构、统一运行入口、测试和复现说明"),
        ("第 8 章 结论", "逐项回答研究问题，说明限制与后续方向"),
        ("附录", "完整参数、补充表格、代码说明和文件清单"),
    ]
    add_table(doc, ["章节", "主要内容"], outline, [5.0, 11.2], "表 10-1 最终课程设计报告主体大纲。正文将以问题—方法—证据—边界形成闭环。")
    add_heading(doc, "10.1 预期最终交付", 2)
    add_bullet(doc, "课程设计报告 DOCX 与 PDF；")
    add_bullet(doc, "四领域原始/处理后数据说明及哈希清单；")
    add_bullet(doc, "统一运行入口、候选模型代码和依赖说明；")
    add_bullet(doc, "表格、图像、模型输出、测试日志和复现报告；")
    add_bullet(doc, "第 10 周阶段汇报 PPT 与最终答辩 PPT。")

    add_heading(doc, "11 初步结论", 1)
    add_body(doc, "第 1—3 周的调研表明，多模态时间序列课程设计的真正难点不是把文本向量拼接进模型，而是证明这种拼接在严格时间条件下具有稳定、可重复的增益。SafeFAME-TS 因此把朴素基线、安全门控、置换检验、时间分段和一次性测试放在方法核心。该路线允许出现三种同样有效的研究结果：多模态候选通过门控；候选因不稳定而回退；数据条件不足而执行安全退化。三种结果都能回答“何时值得融合”这一研究问题。")
    add_body(doc, "当前 Agriculture 与 Climate 两个领域已经完成工程闭环，说明协议可执行；Climate 也显示了审慎解释的必要性：N 的点估计最好，但相对 Last 的区组 Bootstrap 区间跨过 0。后续工作的重点是验收 SocialGood 数据、验证零文本退化契约，并在 Environment 大规模任务上完成最终跨领域检验。")

    add_heading(doc, "参考文献与开源资料", 1)
    refs = [
        "[1] Liu H, et al. Time-MMD: Multi-Domain Multimodal Dataset for Time Series Analysis. NeurIPS 2024 Datasets and Benchmarks Track. https://proceedings.neurips.cc/paper_files/paper/2024/file/8e7768122f3eeec6d77cd2b424b72413-Paper-Datasets_and_Benchmarks_Track.pdf",
        "[2] AdityaLab. Time-MMD official repository. https://github.com/AdityaLab/Time-MMD",
        "[3] Nie Y, et al. A Time Series is Worth 64 Words: Long-term Forecasting with Transformers. ICLR 2023. https://openreview.net/pdf?id=Jbdc0vTOcol",
        "[4] Oreshkin B N, Carpov D, Chapados N, Bengio Y. N-BEATS: Neural basis expansion analysis for interpretable time series forecasting. ICLR 2020. https://openreview.net/forum?id=r1ecqn4YwB",
        "[5] Zeng A, et al. Are Transformers Effective for Time Series Forecasting? AAAI 2023, 37(9). https://doi.org/10.1609/aaai.v37i9.26317",
        "[6] Jin M, et al. Time-LLM: Time Series Forecasting by Reprogramming Large Language Models. ICLR 2024. https://openreview.net/pdf?id=Unb5CVPtae",
        "[7] Reimers N, Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP 2019. https://aclanthology.org/D19-1410/",
        "[8] Wang W, et al. MiniLM: Deep Self-Attention Distillation for Task-Agnostic Compression of Pre-Trained Transformers. NeurIPS 2020. https://proceedings.neurips.cc/paper/2020/hash/3f5ee243547dee91fbd053c1c4a845aa-Abstract.html",
        "[9] Politis D N, Romano J P. The Stationary Bootstrap. Journal of the American Statistical Association, 1994, 89(428): 1303–1313. https://doi.org/10.1080/01621459.1994.10476870",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        set_paragraph(p, first_line=False, line=18, after=4)
        r = p.add_run(ref)
        set_run_font(r, "宋体", 9)

    # Core properties
    props = doc.core_properties
    props.title = "SafeFAME-TS 第1—3周选题调研与初步技术方案"
    props.subject = "人工智能课程设计"
    props.author = "三人课程设计小组"
    props.keywords = "多模态时间序列, 安全门控, 防泄漏, 负迁移"
    props.comments = "成员姓名与学号由小组提交前填写。"

    # Repeat headers for all tables and keep captions close.
    for table in doc.tables:
        set_repeat_table_header(table.rows[0])
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build_doc()
